import sys
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from docx import Document
from docx.shared import Pt,Mm,RGBColor

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from repair_document import repair_copy

class RepairTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)
        doc=Document()
        doc.styles['Normal'].font.size=Pt(14)
        doc.sections[0].page_width=Mm(216);doc.sections[0].page_height=Mm(279)
        r=doc.add_paragraph('Проверка сохранности текста и исправления формата. '*8).runs[0]
        r.font.size=Pt(7);r.font.color.rgb=RGBColor(255,0,0)
        special=doc.add_paragraph('Формула и подстрочный символ. '*4).runs[0]
        special.font.size=Pt(6);special.font.subscript=True
        doc.sections[0].header.paragraphs[0].text='Header must stay unchanged'
        self.source=self.root/'source.docx';doc.save(self.source)

    def test_copy_preserves_source_text_and_other_zip_parts(self):
        raw=self.source.read_bytes()
        original_text=[p.text for p in Document(self.source).paragraphs]
        result=repair_copy(self.source,self.root/'out')
        self.assertEqual(self.source.read_bytes(),raw)
        self.assertTrue(result['source_unchanged'])
        self.assertEqual([p.text for p in Document(result['output_file']).paragraphs],original_text)
        self.assertTrue(Path(result['change_log']).exists())
        with zipfile.ZipFile(self.source) as before,zipfile.ZipFile(result['output_file']) as after:
            self.assertEqual(before.namelist(),after.namelist())
            for name in before.namelist():
                if name!='word/document.xml':self.assertEqual(before.read(name),after.read(name),name)
        self.assertEqual(Document(result['output_file']).paragraphs[1].runs[0].font.size.pt,6)
        self.assertTrue(result['skipped'])
        for code in ('Word.PageSizeNotA4','Word.TextColorNotBlack'):
            self.assertNotIn(code,result['after_rule_ids'])

    def test_dry_run_writes_nothing(self):
        before=set(self.root.rglob('*'))
        result=repair_copy(self.source,dry_run=True)
        self.assertTrue(result['changes'])
        self.assertEqual(set(self.root.rglob('*')),before)

    def test_each_run_creates_new_copy(self):
        one=repair_copy(self.source,self.root/'out')
        two=repair_copy(self.source,self.root/'out')
        self.assertNotEqual(one['output_file'],two['output_file'])

    def test_unvalidated_context_refused(self):
        with self.assertRaises(ValueError):repair_copy(self.source,work_type='specialist')

    def test_master_copy_and_context_evidence(self):
        from unittest.mock import patch
        import repair_document
        result=repair_copy(self.source,self.root/'master',work_type='master')
        self.assertEqual(result['work_type'],'master')
        self.assertTrue(result['source_unchanged'])
        self.assertNotIn('Word.PageSizeNotA4',result['after_rule_ids'])
        registry=repair_document.load_rule_registry()
        repair_document.verification_context(registry['Word.PageSizeNotA4'],'docx','master')['verification_evidence_sha256']='invalid'
        with patch.object(repair_document,'load_rule_registry',return_value=registry):
            with self.assertRaises(ValueError):repair_copy(self.source,work_type='master')

    def test_pdf_refused(self):
        with self.assertRaises(ValueError):repair_copy(self.root/'paper.pdf')

    def test_unknown_action_refused(self):
        with self.assertRaises(ValueError):repair_copy(self.source,actions=['line_spacing'])

    def test_missing_body_target_does_not_guess(self):
        doc=Document(self.source);doc.styles['Normal'].font.size=None;doc.save(self.source)
        result=repair_copy(self.source,actions=['font_size_anomalies'],dry_run=True)
        self.assertEqual(result['changes'],[])
        self.assertTrue(result['skipped'])

    def test_mcp_preview_and_copy_tools(self):
        import subprocess
        requests=[{'jsonrpc':'2.0','id':i,'method':'tools/call','params':{'name':name,'arguments':{'file_path':str(self.source),'output_dir':str(self.root/'mcp')}}}
                  for i,name in enumerate(('preview_thesis_fixes','fix_thesis_copy'),1)]
        run=subprocess.run([sys.executable,str(ROOT/'mcp/server.py')],input='\n'.join(json.dumps(m) for m in requests)+'\n',
                           text=True,encoding='utf-8',capture_output=True,timeout=20)
        self.assertEqual(run.returncode,0,run.stderr)
        outputs=[json.loads(line)['result']['structuredContent'] for line in run.stdout.splitlines()]
        self.assertTrue(outputs[0]['dry_run'])
        self.assertTrue(outputs[1]['source_unchanged'])
        self.assertTrue(Path(outputs[1]['change_log']).exists())

    def test_large_body_font_is_repaired_to_normal_style(self):
        doc=Document(self.source);doc.paragraphs[0].runs[0].font.size=Pt(19);doc.save(self.source)
        result=repair_copy(self.source,actions=['font_size_anomalies'])
        self.assertEqual(Document(result['output_file']).paragraphs[0].runs[0].font.size.pt,14)
