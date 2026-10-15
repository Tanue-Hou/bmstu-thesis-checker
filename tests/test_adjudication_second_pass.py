"""Evidence semantics that prevent unsafe recommendations on complex theses."""
import sys,tempfile,unittest
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(ROOT/'scripts'),str(ROOT/'scripts/testvkr-service')]
from parsers.docx_parser import DocxParser
from issue_locations import locate_finding
from review_policy import assess
from adjudicate_remaining import structural_facts,decide

class SecondPassTests(unittest.TestCase):
    def test_row_locator_is_individual_and_does_not_claim_rendered_height(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'rows.docx';d=Document();t=d.add_table(rows=2,cols=1)
            t.cell(0,0).text='Formula and wrapped table text';h=OxmlElement('w:trHeight');h.set(qn('w:val'),'100');t.rows[0]._tr.get_or_add_trPr().append(h);d.save(path)
            before=path.read_bytes();data=DocxParser(str(path)).parse_all()
            f={'rule_id':'Word.TableRowHeightTooSmall','location':{'table_rows':[{'table':1,'row':1},{'table':1,'row':1}]}}
            loc=locate_finding(f,data,'docx');self.assertEqual(loc['occurrence_count'],1)
            row=loc['occurrences'][0];self.assertEqual(row['kind'],'docx_table_row');self.assertIn('/w:tr[1]',row['xml_path']);self.assertIsNone(row['height_rule_explicit']);self.assertEqual(row['height_measurement'],'specified_xml_property_not_rendered_height');self.assertEqual(path.read_bytes(),before)
            self.assertEqual(assess({**f,'precise_locations':loc},{'document_kind':'vkr'})[0],'layout_property_advisory')
    def test_cover_paragraph_format_candidate_is_not_body_fix(self):
        data={'structure_paragraphs':[{'index':0,'xml_path':'cover','text':'Long cover title'},{'index':1,'xml_path':'intro','text':'ВВЕДЕНИЕ'}],'paragraphs':[{'index':0,'xml_path':'cover','style_name':'Normal','text':'Long cover title'}]}
        f={'rule_id':'Word.AlignmentNotJustified','location':{'paragraphs':[1]}}
        f['precise_locations']=locate_finding(f,data,'docx')
        self.assertEqual(assess(f,{'document_kind':'vkr'})[0],'special_role_requires_review')
    def test_citation_list_and_range_are_retained_without_deleting_sources(self):
        records=[{'index':0,'text':'References [1, 2–4] and [5, с. 21].','xml_path':'p1'},{'index':1,'text':'СПИСОК ЛИТЕРАТУРЫ','xml_path':'p2'},{'index':2,'text':'A numbered source','xml_path':'p3'}]
        facts=structural_facts('Struct.MissingCitations',{'structure_paragraphs':records})
        self.assertEqual(len(facts['citation_candidates']),2)
        self.assertEqual(decide('Struct.MissingCitations',facts,'native')[0],'citation_semantics_requires_review')
    def test_university_word_template_is_not_current_school_truth(self):
        facts=structural_facts('Graduate.UniversityNameInvalid',{'structure_paragraphs':[{'text':'Московский государственный технический университет имени Баумана','index':0,'xml_path':'cover'}]})
        self.assertIn('автономное;автономного',facts['missing_tokens'])
        self.assertEqual(decide('Graduate.UniversityNameInvalid',facts,'native')[1],'template_predicate_not_school_violation')
    def test_empty_parent_direct_text_does_not_mean_empty_child_content(self):
        records=[{'text':'1.1 Конструкторская часть','index':0,'xml_path':'p1'},{'text':'1.1.1 Конструкция','index':1,'xml_path':'p2'},{'text':'Substantial child text','index':2,'xml_path':'p3'}]
        facts=structural_facts('Native.EmptyNumberedSection',{'structure_paragraphs':records},'"1.1 конструкторская часть"')
        self.assertTrue(facts['target_matches'][0]['immediate_child_heading'])
        self.assertEqual(decide('Native.EmptyNumberedSection',facts,'native')[0],'parent_heading_container_advisory')
    def test_same_level_number_collision_is_not_mistaken_for_child_heading(self):
        records=[{'text':'1.1 Конструкторская часть','index':0,'xml_path':'p1'},{'text':'1.1 Ракеты Ариан','index':1,'xml_path':'p2'},{'text':'Substantial section text','index':2,'xml_path':'p3'}]
        facts=structural_facts('Native.EmptyNumberedSection',{'structure_paragraphs':records},'"1.1 конструкторская часть"')
        target=facts['target_matches'][0];self.assertFalse(target['immediate_child_heading']);self.assertTrue(target['consecutive_number_collision'])
        self.assertEqual(decide('Native.EmptyNumberedSection',facts,'native')[0],'confirmed_consecutive_heading_number_collision')

if __name__=='__main__':unittest.main()
