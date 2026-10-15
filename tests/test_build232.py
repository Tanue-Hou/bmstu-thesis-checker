import copy,json,sys,tempfile,unittest,subprocess,os,hashlib
from pathlib import Path
from argparse import Namespace
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(ROOT/'scripts'),str(ROOT/'scripts/testvkr-service')]
from parsers.docx_parser import DocxParser
from checkers.graduate_structure import check_graduate,TITUL
from checkers.entry_document import EntryDocumentChecker
from check_document import check_file
from compare_native_trial import compare,parse_diagnostics,load_capture
from repair_document import repair_copy

def args(build):return Namespace(work_type='master',native_build=build,skip_structure=False,skip_pictures=False,skip_word_format=False,skip_dictionary=False)
class Build232Tests(unittest.TestCase):
    def test_table_cover_and_sdt_toc_document_order(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'cover.docx';d=Document();cell=d.add_table(rows=1,cols=1).cell(0,0);cell.text=' '.join(t.split(';')[0] for t in TITUL)+' ВЫПУСКНАЯ КВАЛИФИКАЦИОННАЯ РАБОТА'
            d.add_paragraph('РЕФЕРАТ');d.add_paragraph('Текст реферата. '*70);d.add_paragraph('СОДЕРЖАНИЕ')
            control=OxmlElement('w:sdt');content=OxmlElement('w:sdtContent');control.append(content)
            for i in range(12):
                p=OxmlElement('w:p');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=f'{i+1} Раздел {i+1}';r.append(t);p.append(r);content.append(p)
            d._element.body.insert(len(d._element.body)-1,control);d.add_paragraph('ВВЕДЕНИЕ');d.add_paragraph('Основной текст. '*100);d.add_paragraph('ЗАКЛЮЧЕНИЕ');d.add_paragraph('СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ');d.add_paragraph('ПРИЛОЖЕНИЕ А');d.save(path)
            data=DocxParser(str(path)).parse_all();codes={f.code for f in check_graduate(data,'master')}
            self.assertTrue(data['structure_paragraphs'][0]['inside_table']);self.assertTrue(any(p['inside_content_control'] for p in data['structure_paragraphs']))
            self.assertNotIn('Graduate.WorkNameInvalid',codes);self.assertNotIn('Graduate.TitlePageMismatch',codes);self.assertNotIn('Graduate.TocEmpty',codes);self.assertFalse(data['native_graduate_scope_supported'])
    def test_table_only_document_not_empty(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'table.docx';d=Document();d.add_table(rows=1,cols=1).cell(0,0).text='Table-only source text';d.save(path)
            data=DocxParser(str(path)).parse_all();self.assertNotIn('Entry.DocumentEmpty',{f.code for f in EntryDocumentChecker(data,'docx','master').check()})
    def test_sdt_heading_location_is_xml_not_wrong_main_paragraph(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'toc.docx';d=Document();control=OxmlElement('w:sdt');content=OxmlElement('w:sdtContent');control.append(content);p=OxmlElement('w:p');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='СОДЕРЖАНИЕ';r.append(t);p.append(r);content.append(p);d._element.body.insert(0,control);d.add_paragraph('ВВЕДЕНИЕ');d.save(path)
            report=check_file(path,args(232));f=next(f for f in report['findings'] if f['rule_id']=='Graduate.TocEmpty');loc=f['precise_locations']['occurrences'][0]
            self.assertEqual(loc['kind'],'docx_xml_paragraph');self.assertIn('sdt',loc['xml_path']);self.assertNotIn('paragraph',loc)
    def test_font_locations_distinguish_front_matter(self):
        with tempfile.TemporaryDirectory() as td:
            from docx.shared import Pt
            path=Path(td)/'roles.docx';d=Document();d.add_paragraph('Министерство науки');d.add_paragraph('Длинный заголовок об исследовании систем '*3).runs[0].font.size=Pt(20);d.add_paragraph('ВВЕДЕНИЕ');d.add_paragraph('Основной текст исследования '*8).runs[0].font.size=Pt(20);d.save(path)
            report=check_file(path,args(232));f=next(f for f in report['findings'] if f['rule_id']=='Word.FontSizeTooLarge');roles={p['content_role'] for p in f['precise_locations']['occurrences']}
            self.assertNotIn('front_matter_candidate_manual_review',roles);self.assertIn('body_or_special_content_unknown',roles)
            self.assertTrue(all(p['native_font_enabled'] for p in f['precise_locations']['occurrences']))
            self.assertEqual({p['paragraph'] for p in f['precise_locations']['occurrences']},{4})
            # The old generic build230 path still exposes front-matter review;
            # native232 suppresses its font count until the exact marker.
            older=next(f for f in check_file(path,args(230))['findings'] if f['rule_id']=='Word.FontSizeTooLarge')
            self.assertIn('front_matter_candidate_manual_review',{p['content_role'] for p in older['precise_locations']['occurrences']})
    def test_native_numbering_ignores_quoted_excerpt_and_preserves_unknown(self):
        text='Файл: example.docx 1. В документе много символов, размер которых превышает 18 пт: "2. Заголовок". 2. Неизвестное сообщение.'
        rows=parse_diagnostics(text,'docx');self.assertEqual(len(rows),2);self.assertEqual(rows[0]['rule_id'],'Word.FontSizeTooLarge');self.assertIsNone(rows[1]['rule_id'])
    def test_build230_proof_does_not_transfer(self):
        p=ROOT/'evidence/dynamic-master/font_7.docx';a=check_file(p,args(230));b=check_file(p,args(232))
        self.assertEqual(next(f for f in a['findings'] if f['rule_id']=='Word.FontSizeTooSmall')['verification_status'],'verified_against_original')
        current=next(f for f in b['findings'] if f['rule_id']=='Word.FontSizeTooSmall');self.assertEqual(current['verification_status'],'verified_against_original');self.assertEqual(current['verification_scope']['build'],232)
        unproven=args(232);unproven.work_type='bachelor';self.assertEqual(next(f for f in check_file(p,unproven)['findings'] if f['rule_id']=='Word.FontSizeTooSmall')['verification_status'],'pending_original_program_comparison')
        self.assertFalse(b['native_reference']['transfer_validated'])
        with self.assertRaises(ValueError):repair_copy(p,dry_run=True,native_build=232)
    def test_real_control_predicates_and_capture_identity(self):
        out=ROOT/'evidence/build-232/controls';report=json.loads((out/'comparison.json').read_text(encoding='utf-8'))
        self.assertTrue(report['all_pass']);self.assertEqual(len(report['cases']),4)
        for case in report['cases']:
            self.assertEqual(hashlib.sha256((out/(case['id']+'.docx')).read_bytes()).hexdigest(),case['sha256'])
            local=json.loads((out/(case['id']+'.local.json')).read_text(encoding='utf-8'));capture=load_capture(out/(case['id']+'.capture.json'));self.assertEqual(compare(local,capture)['build'],232)
        capture['executable_sha256']='bad'
        with self.assertRaises(ValueError):compare(local,capture)
    def test_build_and_screenshot_mismatch_refused(self):
        out=ROOT/'evidence/build-232/controls';local=json.loads((out/'Z01.local.json').read_text(encoding='utf-8'));capture=load_capture(out/'Z01.capture.json')
        local['native_reference']['selected_build']=230
        with self.assertRaises(ValueError):compare(local,capture)
        local['native_reference']['selected_build']=232;capture['screenshots'][0]['sha256']='bad'
        with self.assertRaises(ValueError):compare(local,capture)
    def test_mcp_current_build_and_comparison(self):
        out=ROOT/'evidence/build-232/controls'
        requests=[{'jsonrpc':'2.0','id':1,'method':'tools/call','params':{'name':'list_native_builds'}},{'jsonrpc':'2.0','id':2,'method':'tools/call','params':{'name':'compare_thesis_results','arguments':{'local_report_path':str(out/'Z01.local.json'),'native_capture_path':str(out/'Z01.capture.json')}}},{'jsonrpc':'2.0','id':3,'method':'tools/call','params':{'name':'check_thesis_document','arguments':{'file_path':str(out/'Z01.docx'),'work_type':'master'}}}]
        p=subprocess.run([sys.executable,str(ROOT/'mcp/server.py')],input='\n'.join(json.dumps(r) for r in requests)+'\n',capture_output=True,text=True,encoding='utf-8',env={**os.environ,'PYTHONUTF8':'1'},timeout=30)
        self.assertEqual(p.returncode,0,p.stderr);replies=[json.loads(x)['result']['structuredContent'] for x in p.stdout.splitlines()]
        self.assertEqual(replies[0]['current_build'],232);self.assertEqual(replies[1]['build'],232);self.assertEqual(replies[2]['native_reference']['selected_build'],232)

if __name__=='__main__':unittest.main()
