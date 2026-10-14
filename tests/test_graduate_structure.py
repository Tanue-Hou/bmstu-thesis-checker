import hashlib,json,sys,tempfile,unittest
from argparse import Namespace
from pathlib import Path
from docx import Document
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from check_document import check_file,load_rule_registry

def check(path,mode='master'):
    return check_file(path,Namespace(work_type=mode,skip_structure=False,skip_pictures=False,skip_word_format=False,skip_dictionary=False))

class GraduateTests(unittest.TestCase):
    def test_mcp_check_routes_master_and_phd(self):
        import subprocess
        out=ROOT/'evidence/graduate-structure'
        requests=[{'jsonrpc':'2.0','id':i,'method':'tools/call','params':{'name':'check_thesis_document','arguments':{'file_path':str(out/name),'work_type':mode}}} for i,(name,mode) in enumerate([('g_chars43199.docx','master'),('g_phd_structure.docx','phd')],1)]
        run=subprocess.run([sys.executable,str(ROOT/'mcp/server.py')],input='\n'.join(json.dumps(x) for x in requests)+'\n',capture_output=True,text=True,encoding='utf-8',timeout=20)
        self.assertEqual(run.returncode,0,run.stderr)
        replies=[json.loads(s) for s in run.stdout.splitlines()]
        master=replies[0]['result']['structuredContent']
        phd=replies[1]['result']['structuredContent']
        self.assertIn('Graduate.ThesisSizeTooSmall',{f['rule_id'] for f in master['findings']})
        self.assertEqual(phd['work_type_policy']['toc_heading'],'ОГЛАВЛЕНИЕ')
        finding=next(f for f in phd['findings'] if f['rule_id']=='Graduate.RequiredHeadingsMissing')
        self.assertEqual(finding['verification_scope']['work_type'],'phd')

    def test_master_native_boundary_fixtures(self):
        out=ROOT/'evidence/graduate-structure'
        for name,expected in [('g_toc47',1),('g_toc48',None),('g_chars43199',1),('g_chars43200',None)]:
            findings=[f for f in check(out/(name+'.docx'))['findings'] if f['rule_id']=='Graduate.ThesisSizeTooSmall']
            self.assertEqual(bool(findings),expected is not None)
            if findings:
                self.assertEqual(findings[0]['threshold']['missing_equivalent_pages'],expected)
                self.assertEqual(findings[0]['verification_scope']['work_type'],'master')
        # Correctly formatted native-success fixture must not invent heading-order errors from TOC rows.
        self.assertFalse(check(out/'g_chars43200.docx')['findings'])

    def test_toc_boundaries_and_synthetic_cover(self):
        out=ROOT/'evidence/graduate-structure'
        for name,expected in [('g_cover','Graduate.TocEmpty'),('g_toc_entries5','Graduate.TocSmall'),('g_toc_entries10','Graduate.TocSmall'),('g_toc_entries11',None)]:
            codes={f['rule_id'] for f in check(out/(name+'.docx'))['findings']}
            if expected:self.assertIn(expected,codes)
            else:self.assertFalse({'Graduate.TocEmpty','Graduate.TocSmall'}&codes)
            self.assertNotIn('Graduate.TitlePageMismatch',codes)

    def test_phd_is_a_distinct_mode(self):
        out=ROOT/'evidence/graduate-structure'
        bad=check(out/'g_baseline.docx','phd')['findings']
        self.assertIn('Graduate.WorkNameInvalid',{f['rule_id'] for f in bad})
        wrong=check(out/'g_phd_structure.docx','phd')['findings']
        missing=next(f for f in wrong if f['rule_id']=='Graduate.RequiredHeadingsMissing')
        self.assertEqual(missing['threshold']['missing'],['ОГЛАВЛЕНИЕ'])
        good=check(out/'g_phd_correct_heading.docx','phd')['findings']
        self.assertFalse(any(f['rule_id'].startswith('Graduate.') for f in good))
        self.assertNotIn('Struct.ThesisSizeTooSmall',{f['rule_id'] for f in bad+wrong+good})
        margin=next(f for f in good if f['rule_id']=='Word.LeftMarginInvalid')
        self.assertEqual(margin['threshold']['value'],25)
        self.assertEqual(margin['verification_status'],'verified_against_original')
        self.assertFalse(check(out/'g_phd_left25.docx','phd')['findings'])

    def test_complex_document_does_not_inherit_text_only_verification(self):
        with tempfile.TemporaryDirectory() as td:
            doc=Document(ROOT/'evidence/graduate-structure/g_cover.docx');doc.add_table(rows=1,cols=1)
            path=Path(td)/'complex.docx';doc.save(path)
            findings=[f for f in check(path)['findings'] if f['rule_id'].startswith('Graduate.')]
            self.assertTrue(findings)
            self.assertTrue(all(f['verification_status']!='verified_against_original' for f in findings))

    def test_untested_heading_and_abstract_predicates_remain_pending(self):
        with tempfile.TemporaryDirectory() as td:
            doc=Document(ROOT/'evidence/graduate-structure/g_cover.docx')
            ps=doc.paragraphs;idx=next(i for i,p in enumerate(ps) if p.text=='РЕФЕРАТ')
            ps[idx+1].text='Мало.'
            next(p for p in ps if p.text=='ВВЕДЕНИЕ').text='Другой заголовок'
            path=Path(td)/'untested.docx';doc.save(path)
            findings=[f for f in check(path)['findings'] if f['rule_id'] in {'Graduate.AbstractTooShort','Graduate.RequiredHeadingsMissing'}]
            self.assertEqual(len(findings),2)
            self.assertTrue(all(f['verification_status']=='pending_original_program_comparison' for f in findings))

    def test_evidence_bound_to_each_context(self):
        for rule in load_rule_registry().values():
            if not rule['rule_id'].startswith('Graduate.'):continue
            for mode,context in rule['verification_by_work_type'].items():
                p=ROOT/context['evidence_refs'][0]
                self.assertEqual(hashlib.sha256(p.read_bytes()).hexdigest(),context['verification_evidence_sha256'])
                matrix=json.loads(p.read_text(encoding='utf-8'))
                self.assertTrue(matrix['checks'][rule['rule_id']][mode]['pass'])

if __name__=='__main__':unittest.main()
