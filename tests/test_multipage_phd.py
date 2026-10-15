import copy,hashlib,json,os,subprocess,sys,tempfile,unittest,zipfile
from argparse import Namespace
from pathlib import Path
from unittest.mock import patch
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import fitz
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'scripts'))
from check_document import check_file,load_rule_registry,verification_context
from repair_document import repair_copy
from checkers.structure import StructureChecker
OUT=ROOT/'evidence/multipage-phd-repair'
def check(name,mode='phd'):
    return check_file(OUT/name,Namespace(work_type=mode,skip_structure=False,skip_pictures=False,skip_word_format=False,skip_dictionary=False))
class MultipagePhdTests(unittest.TestCase):
    def test_nonfirst_page_detection_and_locations(self):
        for file,rule,page in [('mp_small.pdf','Word.FontSizeTooSmall',2),('mp_red.pdf','Word.TextColorNotBlack',3),('mp_letter.pdf','Word.PageSizeNotA4',3),('mp_image_above.pdf','Pic.PixelSizeTooLarge',2)]:
            r=check(file);f=next(f for f in r['findings'] if f['rule_id']==rule)
            self.assertEqual(f['verification_status'],'verified_against_original');self.assertTrue(f['verification_evidence_intact'])
            self.assertTrue(f['location'].get('page')==page or page in f['location'].get('pages',[]))
            self.assertTrue(r['native_comparison_limits']['unimplemented_native_checks'])
    def test_pdf_upper_probes_have_real_large_text_and_no_native_warning(self):
        matrix=json.loads((OUT/'comparison.json').read_text(encoding='utf-8'))
        for probe in matrix['pdf_upper_font_probes']:
            with fitz.open(OUT/(probe['case']+'.pdf')) as doc:
                spans=[s for b in doc[1].get_text('dict')['blocks'] if b.get('type')==0 for line in b['lines'] for s in line['spans']]
                self.assertGreater(sum(len(s['text']) for s in spans if s['size']>18),10)
                self.assertTrue(any(abs(s['size']-probe['font_pt'])<.01 for s in spans))
            self.assertFalse(probe['native_upper_warning']);self.assertNotIn('Word.FontSizeTooLarge',{f['rule_id'] for f in check(probe['case']+'.pdf')['findings']})
    def test_pdf_structural_heading_pairs_and_nbsp(self):
        self.assertNotIn('Struct.RequiredHeadingsMissing',{f['rule_id'] for f in check('structure_clean.pdf')['findings']})
        for name,key in [('structure_no_abstract.pdf','abstract'),('structure_wrong_toc.pdf','toc')]:
            f=next(f for f in check(name)['findings'] if f['rule_id']=='Struct.RequiredHeadingsMissing');self.assertEqual(f['threshold']['missing'],[key]);self.assertEqual(f['verification_status'],'verified_against_original')
        spans=[{'text':'Текст содержит реферат и оглавление, но это не заголовки.'},{'text':'ВВЕДЕНИЕ'},{'text':'ЗАКЛЮЧЕНИЕ'},{'text':'СПИСОК\u00a0ИСПОЛЬЗОВАННЫХ\u00a0ИСТОЧНИКОВ'}]
        f=next(f for f in StructureChecker({'spans':spans},'pdf','phd').check() if f.code=='Struct.RequiredHeadingsMissing')
        self.assertEqual(set(f.threshold['missing']),{'abstract','toc'})
    def test_unsupported_pdf_profile_downgrades_evidence(self):
        with tempfile.TemporaryDirectory() as td:
            for rotated,pages in [(True,1),(False,4)]:
                doc=fitz.open();source=fitz.open(OUT/'mp_small.pdf')
                for _ in range(pages):doc.insert_pdf(source,from_page=1,to_page=1)
                if rotated:doc[0].set_rotation(90)
                p=Path(td)/f'{rotated}-{pages}.pdf';doc.save(p);doc.close();source.close()
                r=check_file(p,Namespace(work_type='phd',skip_structure=False,skip_pictures=False,skip_word_format=False,skip_dictionary=False));f=next(f for f in r['findings'] if f['rule_id']=='Word.FontSizeTooSmall')
                self.assertEqual(f['verification_status'],'pending_original_program_comparison')
    def test_doctoral_copy_guards_and_protected_content(self):
        with tempfile.TemporaryDirectory() as td:
            d=Document(OUT/'phd_combined.docx');p=d.add_paragraph();ins=OxmlElement('w:ins');run=OxmlElement('w:r');pr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'14');pr.append(sz);run.append(pr);t=OxmlElement('w:t');t.text='Измененный специальный текст сохраняется.';run.append(t);ins.append(run);p._p.append(ins)
            path=Path(td)/'tracked.docx';d.save(path);raw=path.read_bytes();result=repair_copy(path,Path(td)/'out',body_font_pt=14,work_type='phd');self.assertEqual(path.read_bytes(),raw)
            self.assertTrue(any('tracked change' in s.get('reason','') for s in result['skipped']))
            before=Document(path).element.xpath('//w:ins')[0];after=Document(result['output_file']).element.xpath('//w:ins')[0]
            from lxml import etree
            self.assertEqual(etree.tostring(before),etree.tostring(after))
            d.add_section(WD_SECTION_START.NEW_PAGE);d.save(path)
            with self.assertRaisesRegex(ValueError,'one portrait section'):repair_copy(path,work_type='phd')
    def test_doctoral_repair_evidence_is_required(self):
        import repair_document
        registry=load_rule_registry();verification_context(registry['Word.PageSizeNotA4'],'docx','phd')['repair_evidence_sha256']='bad'
        with patch.object(repair_document,'load_rule_registry',return_value=registry):
            with self.assertRaisesRegex(ValueError,'native before/after evidence'):repair_copy(OUT/'phd_combined.docx',work_type='phd',dry_run=True)
    def test_damaged_check_evidence_is_downgraded(self):
        import check_document
        registry=load_rule_registry();verification_context(registry['Word.FontSizeTooSmall'],'pdf','phd')['verification_evidence_sha256']='bad'
        with patch.object(check_document,'load_rule_registry',return_value=registry):
            f=next(f for f in check('mp_small.pdf')['findings'] if f['rule_id']=='Word.FontSizeTooSmall');self.assertFalse(f['verification_evidence_intact']);self.assertEqual(f['verification_status'],'pending_original_program_comparison')
    def test_mcp_doctoral_preview_and_copy(self):
        with tempfile.TemporaryDirectory() as td:
            requests=[{'jsonrpc':'2.0','id':i,'method':'tools/call','params':{'name':tool,'arguments':{'file_path':str(OUT/'phd_combined.docx'),'output_dir':td,'work_type':'phd','body_font_pt':14}}} for i,tool in enumerate(['preview_thesis_fixes','fix_thesis_copy'],1)]
            run=subprocess.run([sys.executable,str(ROOT/'mcp/server.py')],input='\n'.join(json.dumps(r) for r in requests)+'\n',text=True,encoding='utf-8',capture_output=True,timeout=30,env={**os.environ,'PYTHONUTF8':'1'})
            self.assertEqual(run.returncode,0,run.stderr);results=[json.loads(line)['result']['structuredContent'] for line in run.stdout.splitlines()]
            self.assertEqual(len(results[0]['changes']),4);self.assertTrue(results[0]['dry_run']);self.assertTrue(results[1]['source_unchanged']);self.assertEqual(results[1]['after_rule_ids'],[]);self.assertTrue(any('cover/front matter' in x['reason'] for x in results[0]['skipped']));self.assertTrue(Path(results[1]['change_log']).exists())
if __name__=='__main__':unittest.main()
