import json, subprocess, sys, tempfile, unittest
from pathlib import Path
from docx import Document
ROOT=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(ROOT/'scripts'),str(ROOT/'scripts/testvkr-service')]
from review_policy import annotate
from checkers.graduate_structure import check_graduate

class AgentCliTests(unittest.TestCase):
    def test_oblique_cover_name_and_unrelated_text(self):
        def codes(text):
            return {f.code for f in check_graduate({'paragraphs':[{'text':text}]},'master')}
        self.assertNotIn('Graduate.WorkNameInvalid',codes('К ВЫПУСКНОЙ КВАЛИФИКАЦИОННОЙ РАБОТЕ'))
        self.assertIn('Graduate.WorkNameInvalid',codes('Квалификационный работник выпускного отдела'))
    def test_role_policy_does_not_hide_candidates_or_allow_repairs(self):
        r={'document_kind':'vkr','findings':[{'rule_id':'Word.FontSizeTooLarge','precise_locations':{'occurrences':[{'content_role':'front_matter_candidate_manual_review'}],'truncated':False}}]}
        annotate(r)
        self.assertEqual(len(r['findings']),1)
        self.assertEqual(r['findings'][0]['action_eligibility']['status'],'special_role_requires_review')
        self.assertFalse(r['findings'][0]['action_eligibility']['automatic_repair_allowed'])
        r['findings'][0]['precise_locations']['truncated']=True
        annotate(r)
        self.assertEqual(r['findings'][0]['action_eligibility']['status'],'evidence_required')
    def test_cli_json_check_and_explicit_new_build_scope(self):
        cmd=[sys.executable,str(ROOT/'scripts/thesis_cli.py')]
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'a.docx';d=Document();d.add_paragraph('ordinary content');d.save(p)
            result=subprocess.run(cmd+['check',str(p)],capture_output=True,encoding='utf-8')
            self.assertEqual(result.returncode,0,result.stderr)
            data=json.loads(result.stdout)
            self.assertEqual(data['report']['native_reference']['selected_build'],232)
            self.assertTrue(data['agent_contract']['build232_copy_repair_enabled'])
            self.assertEqual(data['agent_contract']['build232_repair_capability']['actions'],['font_size_anomalies'])
            result=subprocess.run(cmd+['plan',str(p)],capture_output=True,encoding='utf-8')
            self.assertEqual(result.returncode,2)
            self.assertEqual(json.loads(result.stdout)['status'],'error')
    def test_cli_stale_plan_refused(self):
        cmd=[sys.executable,str(ROOT/'scripts/thesis_cli.py')]
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'a.docx';d=Document();d.add_paragraph('ordinary content');d.save(p)
            result=subprocess.run(cmd+['plan',str(p),'--native-build','230','--actions','page_size_a4'],capture_output=True,encoding='utf-8')
            self.assertEqual(result.returncode,0,result.stdout+result.stderr)
            plan=Path(td)/'plan.json';plan.write_text(result.stdout,encoding='utf-8')
            d.add_paragraph('changed');d.save(p)
            result=subprocess.run(cmd+['apply',str(p),'--plan',str(plan),'--output-dir',td],capture_output=True,encoding='utf-8')
            self.assertEqual(result.returncode,2)
            self.assertIn('Source changed',json.loads(result.stdout)['message'])
            self.assertFalse(list(Path(td).glob('*.repaired-*.docx')))

if __name__=='__main__':unittest.main()
