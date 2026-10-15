import json,sys,tempfile,unittest,zipfile
from pathlib import Path
from argparse import Namespace
from docx import Document
from docx.shared import Pt
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'scripts'))
from check_document import check_file,parse_document,verification_context,load_rule_registry
from repair_document import repair_copy
from checkers.word_format import WordFormatChecker
OUT=ROOT/'evidence/format-expansion'
def check(name,mode='phd'):
    return check_file(OUT/name,Namespace(work_type=mode,skip_structure=False,skip_pictures=False,skip_word_format=False,skip_dictionary=False))
class FormatExpansionTests(unittest.TestCase):
    def test_native_left_margin_adjacent_twips(self):
        for twips,expected in [(1405,True),(1406,False),(1428,False),(1429,True)]:
            f=[f for f in check(f'margin_twips{twips}.docx')['findings'] if f['rule_id']=='Word.LeftMarginInvalid']
            self.assertEqual(bool(f),expected)
            if f:self.assertEqual(f[0]['verification_status'],'verified_against_original')
    def test_other_margin_axes(self):
        for axis,target,rule in [('right',10,'Right'),('top',20,'Top'),('bottom',20,'Bottom')]:
            for delta,expected in [('1',False),('3',True)]:
                f=[f for f in check(f'margin_{axis}_{target}_{delta}.docx')['findings'] if f['rule_id']=='Word.'+rule+'MarginInvalid']
                self.assertEqual(bool(f),expected)
                if f:self.assertEqual(f[0]['verification_status'],'verified_against_original')
    def test_table_font_counter_and_image_counter(self):
        self.assertFalse(check('complex_clean.docx')['findings'])
        small=next(f for f in check('complex_table7.docx')['findings'] if f['rule_id']=='Word.FontSizeTooSmall')
        self.assertGreater(small['threshold']['table_count_chars'],10)
        self.assertEqual(small['verification_status'],'verified_against_original')
        image=next(f for f in check('complex_image_above.docx')['findings'] if f['rule_id']=='Pic.PixelSizeTooLarge')
        self.assertEqual(image['verification_status'],'verified_against_original')
    def test_merged_cell_not_counted_twice_and_body_table_aggregate(self):
        with tempfile.TemporaryDirectory() as td:
            d=Document();cell=d.add_table(rows=1,cols=2).cell(0,0);cell.merge(d.tables[0].cell(0,1))
            cell.paragraphs[0].text='Данные';cell.paragraphs[0].runs[0].font.size=Pt(7)
            p=Path(td)/'merged.docx';d.save(p);_,data=parse_document(p)
            self.assertNotIn('Word.FontSizeTooSmall',{f.code for f in WordFormatChecker(data,'docx','phd').check()})
            d.add_paragraph('Данные').runs[0].font.size=Pt(7);d.save(p);_,data=parse_document(p)
            f=next(f for f in WordFormatChecker(data,'docx','phd').check() if f.code=='Word.FontSizeTooSmall')
            self.assertEqual(f.threshold['count_chars'],12);self.assertEqual(f.threshold['table_count_chars'],6)
    def test_pdf_proof_is_not_a_format_work_type_cross_product(self):
        for name,rule in [('pdf_font7.pdf','Word.FontSizeTooSmall'),('pdf_letter.pdf','Word.PageSizeNotA4'),('pdf_red.pdf','Word.TextColorNotBlack')]:
            f=next(f for f in check(name)['findings'] if f['rule_id']==rule)
            self.assertEqual(f['verification_status'],'verified_against_original')
            self.assertEqual(f['verification_scope']['format'],'pdf')
            f=next(f for f in check(name,'master')['findings'] if f['rule_id']==rule)
            self.assertEqual(f['verification_status'],'pending_original_program_comparison')
            self.assertEqual(f['evidence_refs'],[])
        self.assertIsNotNone(verification_context(load_rule_registry()['Word.PageSizeNotA4'],'docx','phd'))
    def test_complex_pdf_and_no_invented_phd_floor(self):
        for name,rule in [('pdf_complex_table7.pdf','Word.FontSizeTooSmall'),('pdf_complex_image_above.pdf','Pic.PixelSizeTooLarge')]:
            findings=check(name)['findings']
            f=next(f for f in findings if f['rule_id']==rule)
            self.assertEqual(f['verification_status'],'verified_against_original')
            self.assertNotIn('Struct.ThesisSizeTooSmall',{f['rule_id'] for f in findings})
        self.assertNotIn('Pic.PixelSizeTooLarge',{f['rule_id'] for f in check('pdf_complex_clean.pdf')['findings']})
    def test_pdf_upper_font_not_promoted(self):
        registry=load_rule_registry()
        self.assertIsNone(verification_context(registry['Word.FontSizeTooLarge'],'pdf','phd'))
        self.assertNotIn('Word.FontSizeTooLarge',{f['rule_id'] for f in check('pdf_font18_5.pdf')['findings']})
    def test_copy_repair_records_table_preservation(self):
        with tempfile.TemporaryDirectory() as td:
            source=OUT/'complex_table7.docx';raw=source.read_bytes()
            result=repair_copy(source,td,actions=['font_size_anomalies'],work_type='master')
            self.assertEqual(source.read_bytes(),raw)
            self.assertIn('Word.FontSizeTooSmall',result['manual_only_rule_ids'])
            self.assertTrue(any('Table text' in s.get('reason','') for s in result['skipped']))
            with zipfile.ZipFile(source) as a,zipfile.ZipFile(result['output_file']) as b:
                self.assertTrue(all(a.read(n)==b.read(n) for n in a.namelist()))
    def test_mcp_pdf_runtime_context(self):
        import subprocess
        request={'jsonrpc':'2.0','id':1,'method':'tools/call','params':{'name':'check_thesis_document','arguments':{'file_path':str(OUT/'pdf_font7.pdf'),'work_type':'phd','native_build':230}}}
        p=subprocess.run([sys.executable,str(ROOT/'mcp/server.py')],input=json.dumps(request)+'\n',capture_output=True,text=True,encoding='utf-8',timeout=20)
        self.assertEqual(p.returncode,0,p.stderr)
        r=json.loads(p.stdout)['result']['structuredContent']
        f=next(f for f in r['findings'] if f['rule_id']=='Word.FontSizeTooSmall')
        self.assertEqual(f['verification_status'],'verified_against_original');self.assertEqual(f['verification_scope']['format'],'pdf')
if __name__=='__main__':unittest.main()
