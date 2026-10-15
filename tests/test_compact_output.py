"""Regression coverage for the compact report projection and public entrypoints."""
import copy
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
TEST_TEMP_ROOT = ROOT / "evidence" / "optimization-audit-r1" / "test-fixtures"
sys.path.insert(0, str(ROOT / "scripts"))
from compact_report import project_report


DECISION_FIELDS = (
    "rule_id", "severity", "verification_status", "verification_scope",
    "verification_by_build", "threshold_origin", "threshold", "confidence",
    "fix_priority", "action_eligibility",
)


def make_synthetic_docx(path, paragraphs=100):
    doc = Document()
    for index in range(paragraphs):
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(
            f"Synthetic compact report fixture {index:03d}. "
            "Repeated neutral formatting sample for output projection."
        )
        run.font.size = Pt(24)
    doc.save(path)


def temporary_directory():
    TEST_TEMP_ROOT.mkdir(parents=True, exist_ok=True)
    return tempfile.TemporaryDirectory(dir=TEST_TEMP_ROOT)


def decision_signature(report):
    return [
        {field: finding.get(field) for field in DECISION_FIELDS}
        for finding in report.get("findings", [])
    ]


def location_count_total(report):
    total = 0
    for finding in report.get("findings", []):
        locations = finding.get("precise_locations") or {}
        total += locations.get("occurrence_count", len(locations.get("occurrences", [])))
    return total


def serialized_size(value, *, compact=False):
    kwargs = {"ensure_ascii": False}
    if compact:
        kwargs["separators"] = (",", ":")
    else:
        kwargs["indent"] = 2
    return len(json.dumps(value, **kwargs).encode("utf-8"))


def run_stdio_mcp(messages):
    payload = "\n".join(json.dumps(item, ensure_ascii=False) for item in messages) + "\n"
    process = subprocess.run(
        [sys.executable, str(ROOT / "mcp" / "server.py")],
        input=payload, capture_output=True, text=True, encoding="utf-8", timeout=90,
    )
    if process.returncode:
        raise AssertionError(f"MCP process failed ({process.returncode}): {process.stderr[-3000:]}")
    replies = [json.loads(line) for line in process.stdout.splitlines() if line.strip()]
    by_id = {item.get("id"): item for item in replies}
    for message in messages:
        if "id" in message and (message["id"] not in by_id or "error" in by_id[message["id"]]):
            raise AssertionError(f"MCP request failed: {by_id.get(message.get('id'))}")
    return by_id


class CompactOutputTests(unittest.TestCase):
    def test_projection_is_non_mutating_and_preserves_decisions_and_location_totals(self):
        original = {
            "summary": {"total": 1, "major": 1},
            "native_reference": {"selected_build": 232, "transfer_validated": False},
            "findings": [{
                "rule_id": "Word.FontSizeTooLarge",
                "severity": "major",
                "verification_status": "verified_against_original",
                "verification_scope": {"format": "docx", "work_type": "master"},
                "verification_by_build": {"232": {"docx:master": {"status": "verified"}}},
                "threshold_origin": "native232_context",
                "threshold": {"value": 18, "unit": "pt", "occurrence_count": 5},
                "confidence": "high",
                "fix_priority": "manual_review",
                "action_eligibility": {"status": "not_eligible", "automatic_repair_allowed": False},
                "evidence_refs": [f"evidence/ref-{index}.json" for index in range(5)],
                "precise_locations": {
                    "occurrences": [{"paragraph": index + 1, "xml_path": f"/p[{index + 1}]"}
                                     for index in range(5)],
                    "occurrence_count": 5,
                    "truncated": False,
                    "limit": 100,
                },
            }],
            "native_paragraph_observations": {
                "native_paragraph_records": [{"index": index} for index in range(7)],
                "records_truncated": False,
            },
        }
        before = copy.deepcopy(original)

        compact = project_report(original, "compact")

        self.assertEqual(original, before, "projection must not mutate the caller's report")
        self.assertEqual(compact["output_detail"]["level"], "compact")
        self.assertTrue(compact["output_detail"]["complete_findings_list"])
        self.assertEqual(len(compact["findings"]), len(original["findings"]))
        source_finding, compact_finding = original["findings"][0], compact["findings"][0]
        for field in DECISION_FIELDS:
            self.assertEqual(compact_finding.get(field), source_finding.get(field), field)
        self.assertEqual(compact_finding["precise_locations"]["occurrence_count"], 5)
        self.assertEqual(len(compact_finding["precise_locations"]["occurrences"]), 3)
        self.assertTrue(compact_finding["precise_locations"]["truncated"])
        self.assertEqual(compact_finding["precise_locations"]["limit"], 3)
        self.assertEqual(compact_finding["evidence_ref_count"], 5)
        self.assertEqual(compact_finding["evidence_refs"], source_finding["evidence_refs"][:1])
        self.assertTrue(compact_finding["evidence_refs_truncated"])
        self.assertEqual(location_count_total(compact), location_count_total(original))
        self.assertEqual(len(compact["native_paragraph_observations"]["native_paragraph_records"]), 3)
        self.assertTrue(compact["native_paragraph_observations"]["records_truncated"])

    def test_compact_cli_matches_full_decisions_and_review_rejects_compact(self):
        with temporary_directory() as temporary:
            temp = Path(temporary)
            source = temp / "synthetic.docx"
            make_synthetic_docx(source)
            cli = [sys.executable, str(ROOT / "scripts" / "thesis_cli.py")]

            def check(detail):
                run = subprocess.run(
                    cli + ["check", str(source), "--work-type", "master", "--native-build", "230",
                           "--detail", detail],
                    capture_output=True, text=True, encoding="utf-8", timeout=90,
                )
                self.assertEqual(run.returncode, 0, run.stderr + run.stdout[-3000:])
                return json.loads(run.stdout)

            full, compact = check("full"), check("compact")
            full_report, compact_report = full["report"], compact["report"]
            self.assertEqual(full_report["source_sha256"], compact_report["source_sha256"])
            self.assertEqual(full_report["native_reference"], compact_report["native_reference"])
            self.assertEqual(full_report["summary"], compact_report["summary"])
            self.assertEqual(decision_signature(full_report), decision_signature(compact_report))
            self.assertEqual(location_count_total(full_report), location_count_total(compact_report))
            self.assertEqual(full["agent_contract"], compact["agent_contract"])
            self.assertLess(serialized_size(compact), serialized_size(full))
            self.assertEqual(compact_report["output_detail"]["level"], "compact")

            compact_path = temp / "compact.json"
            compact_path.write_text(json.dumps(compact, ensure_ascii=False), encoding="utf-8")
            review = subprocess.run(
                cli + ["review", str(compact_path)], capture_output=True,
                text=True, encoding="utf-8", timeout=30,
            )
            self.assertEqual(review.returncode, 2)
            self.assertIn("Review needs a full report", json.loads(review.stdout)["message"])

    def test_stdio_mcp_compact_preserves_document_and_folder_results(self):
        with temporary_directory() as temporary:
            temp = Path(temporary)
            source = temp / "synthetic.docx"
            make_synthetic_docx(source)
            common = {"file_path": str(source), "work_type": "master", "native_build": 230}
            folder_common = {"folder_path": str(temp), "recursive": False,
                             "work_type": "master", "native_build": 230}
            messages = [
                {"jsonrpc": "2.0", "id": 1, "method": "initialize", "params": {}},
                {"jsonrpc": "2.0", "id": 2, "method": "tools/call", "params": {
                    "name": "check_thesis_document", "arguments": {**common, "detail": "full"}}},
                {"jsonrpc": "2.0", "id": 3, "method": "tools/call", "params": {
                    "name": "check_thesis_document", "arguments": {**common, "detail": "compact"}}},
                {"jsonrpc": "2.0", "id": 4, "method": "tools/call", "params": {
                    "name": "check_thesis_folder", "arguments": {**folder_common, "detail": "full"}}},
                {"jsonrpc": "2.0", "id": 5, "method": "tools/call", "params": {
                    "name": "check_thesis_folder", "arguments": {**folder_common, "detail": "compact"}}},
            ]
            replies = run_stdio_mcp(messages)
            full_doc = replies[2]["result"]["structuredContent"]
            compact_doc = replies[3]["result"]["structuredContent"]
            self.assertEqual(full_doc["source_sha256"], compact_doc["source_sha256"])
            self.assertEqual(full_doc["native_reference"], compact_doc["native_reference"])
            self.assertEqual(full_doc["summary"], compact_doc["summary"])
            self.assertEqual(decision_signature(full_doc), decision_signature(compact_doc))
            self.assertEqual(location_count_total(full_doc), location_count_total(compact_doc))
            self.assertLess(serialized_size(compact_doc, compact=True),
                            serialized_size(full_doc, compact=True))

            full_folder = replies[4]["result"]["structuredContent"]
            compact_folder = replies[5]["result"]["structuredContent"]
            self.assertEqual(full_folder["summary"], compact_folder["summary"])
            self.assertEqual(len(full_folder["results"]), len(compact_folder["results"]))
            self.assertEqual(full_folder["results"][0]["source_sha256"],
                             compact_folder["results"][0]["source_sha256"])
            self.assertEqual(decision_signature(full_folder["results"][0]),
                             decision_signature(compact_folder["results"][0]))
            self.assertEqual(location_count_total(full_folder["results"][0]),
                             location_count_total(compact_folder["results"][0]))
            self.assertLess(serialized_size(compact_folder, compact=True),
                            serialized_size(full_folder, compact=True))


if __name__ == "__main__":
    unittest.main()
