import copy,hashlib,json,sys,tempfile,unittest,importlib.util
from pathlib import Path
from argparse import Namespace
from docx import Document
from docx.shared import Pt
ROOT=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(ROOT/'scripts'),str(ROOT/'scripts'/'testvkr-service')]
from school_templates import load_template,validate_template,family,iter_text
from check_document import check_file
from issue_locations import locate_finding
from repair_document import repair_copy
from checkers.structure import StructureChecker

class SchoolLocationTests(unittest.TestCase):
    def test_profiles_and_scope(self):
        a=load_template('bmstu-master-reference',work_type='master')
        b=load_template('bmstu-dissertation-reference',work_type='phd')
        self.assertEqual(a['targets']['margins_mm']['bottom'],20)
        self.assertEqual(b['targets']['margins_mm']['bottom'],25)
        self.assertFalse(a['applicability_confirmed']);self.assertFalse(b['auto_repair'])
        with self.assertRaises(ValueError):load_template('../escape')
        with self.assertRaises(ValueError):load_template(a['id'],work_type='phd')
        with self.assertRaises(ValueError):load_template(a['id'],path='anything')
    def test_invalid_targets_and_provenance(self):
        base=load_template('bmstu-master-reference')
        for change in [lambda x:x.update(sources=[]),lambda x:x['targets'].update(body_font_pt=float('nan')),lambda x:x['targets'].update(unknown=1),lambda x:x.update(auto_repair=True),lambda x:x['targets'].update(margins_mm={'left':True}),lambda x:x['sources'][0].update(url='file:///tmp/a')]:
            value=copy.deepcopy(base);change(value)
            with self.assertRaises(ValueError):validate_template(value)
    def test_font_postscript_variants(self):
        for s in ['ABCDEF+TimesNewRomanPS-BoldMT','TimesNewRomanPS-ItalicMT','Times New Roman']:
            self.assertEqual(family(s),family('Times New Roman'))
    def test_pdf_locations_bound_and_stable(self):
        data={'spans':[{'text':'abc','font_size_pt':6,'page_index':1,'bbox':[10,20,30,40],'block_index':3,'line_index':4,'span_index':0}]*3}
        f={'rule_id':'Word.FontSizeTooSmall','threshold':{'value':8}}
        a=locate_finding(f,data,'pdf',limit=2)
        self.assertEqual(a['occurrence_count'],3);self.assertTrue(a['truncated'])
        self.assertEqual(a['occurrences'][0]['page'],2)
        self.assertEqual(a,locate_finding(f,data,'pdf',limit=2))
        self.assertIn('上下标',a['occurrences'][0]['context_note'])
    def test_table_merged_cell_dedup(self):
        p={'xml_path':'/document/body/tbl/tr/tc/p','runs':[{'text':'x'}]}
        data={'tables':[{'rows':[{'cells':[{'paragraphs':[p]},{'paragraphs':[p]}]}]}]}
        self.assertEqual(len(list(iter_text(data,'docx'))),1)
    def test_dissertation_heading_aliases_and_split_line(self):
        data={'document_kind':'dissertation','spans':[{'text':x,'line_text':line} for x,line in [('ВВЕДЕНИЕ','ВВЕДЕНИЕ'),('ОБЩИЕ','ОБЩИЕ ВЫВОДЫ И ЗАКЛЮЧЕНИЕ'),('ЛИТЕРАТУРА','ЛИТЕРАТУРА')]]}
        findings=[];StructureChecker(data,'pdf','phd')._check_pdf_structure(findings)
        self.assertNotIn('Struct.RequiredHeadingsMissing',[f.code for f in findings])
        data['document_kind']='nkr';findings=[];StructureChecker(data,'pdf','phd')._check_pdf_structure(findings)
        missing=next(f for f in findings if f.code=='Struct.RequiredHeadingsMissing').threshold['missing']
        self.assertIn('abstract',missing);self.assertIn('conclusion',missing)
    def test_template_status_no_native_promotion_and_docx_coordinates(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'sample.docx';d=Document();d.add_paragraph('Обычный основной текст для проверки. '*6).runs[0].font.size=Pt(10);d.save(p)
            args=Namespace(work_type='master',template_id='bmstu-master-reference',document_kind='auto',skip_structure=True,skip_pictures=True,skip_word_format=False,skip_dictionary=True)
            r=check_file(p,args);f=next(f for f in r['findings'] if f['rule_id']=='Template.BodyFontSize')
            self.assertEqual(f['verification_status'],'school_template_reference_not_native_proof')
            self.assertEqual(f['fix_priority'],'manual_review')
            loc=f['precise_locations']['occurrences'][0];self.assertEqual(loc['paragraph'],1);self.assertEqual(loc['run'],1);self.assertTrue(loc['xml_path']);self.assertNotIn('page',loc)
            args.document_kind='dissertation'
            with self.assertRaises(ValueError):check_file(p,args)
    def test_preview_hash_and_plan_identity(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'sample.docx';d=Document();d.styles['Normal'].font.size=Pt(14);d.add_paragraph('Основной текст. '*8).runs[0].font.size=Pt(7);d.save(p)
            preview=repair_copy(p,dry_run=True)
            result=repair_copy(p,Path(td)/'copies',expected_input_sha256=preview['input_sha256'])
            self.assertEqual(preview['plan_id'],result['plan_id']);self.assertTrue(result['source_unchanged'])
            with self.assertRaises(ValueError):repair_copy(p,dry_run=True,expected_input_sha256='0'*64)
    def test_dissertation_repair_guard(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'sample.docx';d=Document();d.add_paragraph('Диссертация на соискание ученой степени кандидата технических наук');d.save(p)
            with self.assertRaises(ValueError):repair_copy(p,work_type='phd',dry_run=True)
    def test_cover_title_preserved(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'sample.docx';d=Document();d.styles['Normal'].font.size=Pt(14)
            d.add_paragraph('Министерство науки и высшего образования')
            d.add_paragraph('Длинное название выпускной квалификационной работы '*3).runs[0].font.size=Pt(20)
            d.add_paragraph('ВВЕДЕНИЕ')
            d.add_paragraph('Обычный основной текст. '*8).runs[0].font.size=Pt(7);d.save(p)
            r=repair_copy(p,dry_run=True)
            self.assertFalse(any(c.get('paragraph')==2 for c in r['changes']))
            self.assertTrue(any(c.get('paragraph')==4 for c in r['changes']))
    def test_native_capture_identity_and_integrity(self):
        from compare_native_trial import compare
        local={'source_sha256':'abc','work_type':'master','document_kind':'vkr','findings':[]}
        self.assertEqual(compare(local,{'status':'pending'})['status'],'pending_native_capture')
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'capture.txt';p.write_text('размер которых меньше 8 пт',encoding='utf-8')
            capture={'status':'captured','input_sha256':'abc','work_type':'master','build':230,'executable_sha256':'0bb43cebc4c7ee196913d2ea2240f7bc53aaac1e6b8f6394e6201024a108b8cb','text_path':str(p),'text_sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
            result=compare(local,capture);self.assertEqual(result['pairs'][0]['status'],'difference_requires_review')
            capture['input_sha256']='wrong'
            with self.assertRaises(ValueError):compare(local,capture)
            capture['input_sha256']='abc';p.write_text('changed',encoding='utf-8')
            with self.assertRaises(ValueError):compare(local,capture)
    def test_mcp_template_discovery_and_arguments(self):
        spec=importlib.util.spec_from_file_location('school_mcp',ROOT/'mcp/server.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
        result=m.handle_tool_call({'name':'list_school_templates'})
        self.assertEqual(len(result['structuredContent']['templates']),2)
        self.assertEqual(m.make_args({'work_type':'master','template_id':'bmstu-master-reference'}).template_id,'bmstu-master-reference')

if __name__=='__main__':unittest.main()
