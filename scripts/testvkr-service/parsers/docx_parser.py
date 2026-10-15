import os
import zipfile
from typing import Dict, Any, List, Optional
from docx import Document
from docx.shared import Inches, Pt, Cm, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from copy import deepcopy
from docx.text.paragraph import Paragraph
from .style_resolver import StyleResolver
from .drawing_geometry import native_extent,vml_textboxes

# OpenXML Namespaces
NS = {
    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'
}

def get_image_dimensions(data: bytes):
    # PNG
    if data[:8] == b'\x89PNG\r\n\x1a\n' and data[12:16] == b'IHDR':
        import struct
        try:
            w, h = struct.unpack('>II', data[16:24])
            return w, h
        except Exception:
            pass
    # JPEG
    elif data[:2] == b'\xff\xd8':
        import struct
        try:
            offset = 2
            while offset < len(data):
                while offset < len(data) and data[offset] != 0xff:
                    offset += 1
                if offset >= len(data):
                    break
                while offset < len(data) and data[offset] == 0xff:
                    offset += 1
                if offset >= len(data):
                    break
                marker = data[offset]
                offset += 1
                if marker in [0xd8, 0xd9, 0x00]:
                    continue
                if 0xd0 <= marker <= 0xd7:
                    continue
                if offset + 2 > len(data):
                    break
                length = struct.unpack('>H', data[offset:offset+2])[0]
                if (0xc0 <= marker <= 0xc3 or 0xc5 <= marker <= 0xc7 or
                    0xc9 <= marker <= 0xcb or 0xcd <= marker <= 0xcf):
                    if offset + 2 + 5 <= len(data):
                        h, w = struct.unpack('>HH', data[offset+3:offset+7])
                        return w, h
                offset += length
        except Exception:
            pass
    # GIF
    elif data[:6] in [b'GIF87a', b'GIF89a'] and len(data) >= 10:
        import struct
        try:
            w, h = struct.unpack('<HH', data[6:10])
            return w, h
        except Exception:
            pass
    # BMP
    elif data[:2] == b'BM' and len(data) >= 26:
        import struct
        try:
            w, h = struct.unpack('<ii', data[18:26])
            return w, abs(h)
        except Exception:
            pass
    return None

class DocxParser:
    def __init__(self, file_path: str):
        self.file_path = file_path
        self.zip_image_headers={}
        with zipfile.ZipFile(file_path) as archive,open(file_path,'rb') as stream:
            stream.seek(max(0,os.path.getsize(file_path)-70000));tail=stream.read()
            zip64_directory=b'PK\x06\x06' in tail or b'PK\x06\x07' in tail
            def zip64_extra(extra):
                offset=0
                while offset+4<=len(extra):
                    tag=int.from_bytes(extra[offset:offset+2],'little');size=int.from_bytes(extra[offset+2:offset+4],'little')
                    if tag==1 or offset+4+size>len(extra):return True
                    offset+=4+size
                return offset!=len(extra)
            for info in archive.infolist():
                if not info.filename.startswith('word/media/'):continue
                stream.seek(info.header_offset);header=stream.read(30)
                if len(header)==30 and header[:4]==b'PK\x03\x04':
                    name_size=int.from_bytes(header[26:28],'little');extra_size=int.from_bytes(header[28:30],'little');stream.read(name_size);extra=stream.read(extra_size)
                    self.zip_image_headers[info.filename]={'zip_local_uncompressed_size':int.from_bytes(header[22:26],'little'),'zip_central_uncompressed_size':info.file_size,'zip_local_flags':int.from_bytes(header[6:8],'little'),'zip_compression_method':int.from_bytes(header[8:10],'little'),'zip_uses_zip64':zip64_directory or int.from_bytes(header[4:6],'little')>=45 or zip64_extra(extra) or zip64_extra(info.extra),'zip_entry':info.filename,'zip_entry_ambiguous':info.filename in self.zip_image_headers}
        self.doc = Document(file_path)
        self.style_resolver = StyleResolver(self.doc)
        self.metadata = {}
        self.parse_basic_metadata()

    def parse_basic_metadata(self):
        # File Size
        self.metadata['file_size_mb'] = os.path.getsize(self.file_path) / (1024 * 1024)

        # Sections (Margins, Page size, Header/Footer)
        self.metadata['sections'] = []
        for sec in self.doc.sections:
            sec_info = {
                'page_width_mm': sec.page_width.mm if sec.page_width is not None else None,
                'page_height_mm': sec.page_height.mm if sec.page_height is not None else None,
                'margin_left_mm': sec.left_margin.mm if sec.left_margin is not None else None,
                'margin_right_mm': sec.right_margin.mm if sec.right_margin is not None else None,
                'margin_top_mm': sec.top_margin.mm if sec.top_margin is not None else None,
                'margin_bottom_mm': sec.bottom_margin.mm if sec.bottom_margin is not None else None,
                'header_distance_cm': sec.header_distance.cm if sec.header_distance is not None else None,
                'footer_distance_cm': sec.footer_distance.cm if sec.footer_distance is not None else None,
                'gutter_mm': sec.gutter.mm if sec.gutter is not None else 0.0,
            }
            self.metadata['sections'].append(sec_info)

    def extract_runs_formatting(self, p_element) -> List[Dict[str, Any]]:
        """Extract formatting for each run in a paragraph, including low-level XML attributes like scale."""
        runs_info = []
        # Find all w:r elements in the paragraph XML
        w_r_list = p_element.findall('.//w:r', NS)

        for w_r in w_r_list:
            text_nodes = w_r.findall('.//w:t', NS)
            text = "".join([t.text for t in text_nodes if t.text])
            if not text:
                continue

            rPr, formatting_sources = self.style_resolver.properties(p_element, w_r)

            # Default properties
            font_size = None
            font_name = None
            color_hex = None
            bg_color_hex = None
            scale = 100.0  # Percentage
            spacing = 0.0  # pt expansion/compression
            is_hidden = False

            if rPr is not None:
                # Font size (w:sz, value is half-points)
                sz = rPr.find('w:sz', NS)
                if sz is not None:
                    val = sz.get(qn('w:val'))
                    if val:
                        font_size = float(val) / 2.0

                # Font family
                rFonts = rPr.find('w:rFonts', NS)
                if rFonts is not None:
                    font_name = rFonts.get(qn('w:ascii'))

                # Color (w:color)
                color = rPr.find('w:color', NS)
                if color is not None:
                    color_hex = color.get(qn('w:val'))

                # Background / shading (w:shd)
                shd = rPr.find('w:shd', NS)
                if shd is not None:
                    bg_color_hex = shd.get(qn('w:fill'))

                # Font Scale (w:w)
                w_w = rPr.find('w:w', NS)
                if w_w is not None:
                    val = w_w.get(qn('w:val'))
                    if val:
                        scale = float(val)

                # Spacing (w:spacing, value is twentieths of a point)
                sp = rPr.find('w:spacing', NS)
                if sp is not None:
                    val = sp.get(qn('w:val'))
                    if val:
                        spacing = float(val) / 20.0

                # Hidden text (w:vanish)
                vanish = rPr.find('w:vanish', NS)
                if vanish is not None:
                    is_hidden = vanish.get(qn('w:val'), 'true') not in ('0', 'false', 'off')

            runs_info.append({
                'text': text,
                'font_size_pt': font_size,
                'font_name': font_name,
                'color_hex': color_hex,
                'bg_color_hex': bg_color_hex,
                'scale_pct': scale,
                'spacing_pt': spacing,
                'is_hidden': is_hidden,
                'formatting_sources': formatting_sources
            })
        return runs_info

    def effective_paragraph(self, paragraph):
        props, sources = self.style_resolver.properties(paragraph._p)
        proxy = Paragraph(deepcopy(paragraph._p), paragraph._parent)
        original = proxy._p.find(qn('w:pPr'))
        if original is not None:
            proxy._p.remove(original)
        proxy._p.insert(0, props)
        return proxy, sources

    def get_paragraph_alignment(self, p):
        return {WD_ALIGN_PARAGRAPH.JUSTIFY: 'justify',
                WD_ALIGN_PARAGRAPH.CENTER: 'center',
                WD_ALIGN_PARAGRAPH.RIGHT: 'right',
                WD_ALIGN_PARAGRAPH.LEFT: 'left'}.get(p.alignment, 'left')

    def parse_paragraphs(self) -> List[Dict[str, Any]]:
        paragraphs_info = []
        for i, p in enumerate(self.doc.paragraphs):
            effective_p, formatting_sources = self.effective_paragraph(p)
            p_format = effective_p.paragraph_format

            # Line Spacing
            # line_spacing = None means default. If it is a multiplier, line_spacing is float (e.g. 1.5)
            # if it's exact spacing, it is a Pt value
            line_spacing = p_format.line_spacing
            line_spacing_type = 'multiplier'
            if isinstance(line_spacing, int):
                line_spacing_type = 'pt'
                line_spacing = line_spacing.pt
            elif line_spacing is not None:
                line_spacing = float(line_spacing)
            else:
                line_spacing = 1.0 # Default fallback

            first_line_indent = p_format.first_line_indent.cm if p_format.first_line_indent else 0.0
            left_indent = p_format.left_indent.cm if p_format.left_indent else 0.0
            right_indent = p_format.right_indent.cm if p_format.right_indent else 0.0
            space_before = p_format.space_before.pt if p_format.space_before else 0.0
            space_after = p_format.space_after.pt if p_format.space_after else 0.0

            alignment = self.get_paragraph_alignment(effective_p)
            text = p.text

            # Extract run details
            runs = self.extract_runs_formatting(p._p)

            paragraphs_info.append({
                'index': i,
                'xml_path':p._p.getroottree().getpath(p._p),
                'style_name':p.style.name,
                'special_content':bool(p._p.xpath('.//w:drawing | .//w:pict | .//w:fldChar | .//w:fldSimple | .//w:instrText | .//m:oMath | .//w:ins | .//w:del | .//w:sdt | .//w:hyperlink')),
                'text': text,
                'line_spacing': line_spacing,
                'line_spacing_type': line_spacing_type,
                'first_line_indent_cm': first_line_indent,
                'left_indent_cm': left_indent,
                'right_indent_cm': right_indent,
                'space_before_pt': space_before,
                'space_after_pt': space_after,
                'alignment': alignment,
                'runs': runs,
                'formatting_sources': formatting_sources
            })
        return paragraphs_info

    def parse_tables(self) -> List[Dict[str, Any]]:
        tables_info = []
        for table_idx, t in enumerate(self.doc.tables):
            rows_info = []

            # Check wrapping/alignment of table in XML
            tblPr = t._tbl.tblPr
            tblpPr = tblPr.find(qn('w:tblpPr'))
            has_wrapping = tblpPr is not None # If w:tblpPr exists, table is floating/wrapped

            for row_idx, r in enumerate(t.rows):
                # Row height extraction
                trPr = r._tr.get_or_add_trPr()
                trHeight = trPr.find(qn('w:trHeight'))
                height_val = None
                height_rule = 'auto'
                if trHeight is not None:
                    val = trHeight.get(qn('w:val'))
                    if val:
                        height_val = float(val) / 20.0 # Twips to points (1 pt = 20 twips)
                        # height in mm: height_val * 0.352778
                        height_val = height_val * 0.352778
                    hRule = trHeight.get(qn('w:hRule'))
                    if hRule:
                        height_rule = hRule # 'exact', 'atLeast', 'auto'

                cells_info = []
                for cell_idx, c in enumerate(r.cells):
                    # Shading (bg color)
                    tcPr = c._tc.get_or_add_tcPr()
                    shd = tcPr.find(qn('w:shd'))
                    bg_color = None
                    if shd is not None:
                        bg_color = shd.get(qn('w:fill'))

                    # Extract paragraphs in this cell
                    cell_paragraphs = []
                    for cp in c.paragraphs:
                        # Extract formatting
                        effective_cp, cell_sources = self.effective_paragraph(cp)
                        line_spacing = effective_cp.paragraph_format.line_spacing
                        line_spacing_type = 'multiplier'
                        if isinstance(line_spacing, int):
                            line_spacing_type = 'pt'
                            line_spacing = line_spacing.pt
                        elif line_spacing is not None:
                            line_spacing = float(line_spacing)
                        else:
                            line_spacing = 1.0

                        runs = self.extract_runs_formatting(cp._p)
                        cell_paragraphs.append({
                            'xml_path': cp._p.getroottree().getpath(cp._p),
                            'text': cp.text,
                            'line_spacing': line_spacing,
                            'line_spacing_type': line_spacing_type,
                            'formatting_sources': cell_sources,
                            'runs': runs
                        })

                    cells_info.append({
                        'index': cell_idx,
                        'bg_color_hex': bg_color,
                        'paragraphs': cell_paragraphs
                    })

                rows_info.append({
                    'index': row_idx,
                    'xml_path': r._tr.getroottree().getpath(r._tr),
                    'height_mm': height_val,
                    'height_rule': height_rule,
                    'height_rule_explicit': trHeight.get(qn('w:hRule')) if trHeight is not None else None,
                    'height_measurement': 'specified_xml_property_not_rendered_height',
                    'cells': cells_info
                })

            tables_info.append({
                'index': table_idx,
                'has_wrapping': has_wrapping,
                'rows': rows_info
            })
        return tables_info

    def parse_drawings_and_images(self) -> List[Dict[str, Any]]:
        images_info = []
        # Find all w:drawing elements in the document XML
        doc_element = self.doc.element.body
        drawings = doc_element.findall('.//w:drawing', NS)

        # Get document relationships
        rels = {}
        try:
            rels = self.doc.part.rels
        except Exception:
            pass

        for idx, d in enumerate(drawings):
            # Check if it is floating (wp:anchor) or inline (wp:inline)
            anchor = d.find('wp:anchor', NS)
            inline = d.find('wp:inline', NS)

            placement = 'inline'
            extent = None
            if anchor is not None:
                placement = 'floating'
                extent = anchor.find('wp:extent', NS)
            elif inline is not None:
                extent = inline.find('wp:extent', NS)

            width_px = None
            height_px = None
            native_width=native_height=native_area=None
            wrap_mode=None
            if anchor is not None:
                wrap_mode=next((etree.QName(node).localname for node in anchor if etree.QName(node).localname.startswith('wrap')),None)
            if extent is not None:
                # Size is in EMUs (English Metric Units). 1 inch = 914400 EMUs, 1 px ≈ 9525 EMUs (at 96 dpi)
                cx = extent.get('cx')
                cy = extent.get('cy')
                if cx and cy:
                    native_width,native_height,native_area=native_extent(cx,cy)
                    width_px = float(cx) / 9525.0
                    height_px = float(cy) / 9525.0

            byte_size_mb = 0.0
            image_extension = None
            png_bit_depth = png_color_type = None
            png_structure_verified = False
            zip_metadata={}
            display_width_px, display_height_px = width_px, height_px
            # Rendered extent at 96dpi is not the image's native pixel count.
            width_px, height_px = None, None

            # Find image relationship and extract raw dimensions/bytes size
            blip = d.find('.//a:blip', NS)
            if blip is not None:
                embed_rid = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                if embed_rid and embed_rid in rels:
                    try:
                        rel = rels[embed_rid]
                        zip_metadata=self.zip_image_headers.get(str(rel.target_part.partname).lstrip('/'),{})
                        blob = rel.target_part.blob
                        image_extension = 'png' if blob.startswith(b'\x89PNG\r\n\x1a\n') else 'other'
                        if image_extension=='png' and len(blob)>=26:png_bit_depth,png_color_type=blob[24],blob[25]
                        if image_extension=='png':
                            try:
                                from io import BytesIO
                                from PIL import Image
                                with Image.open(BytesIO(blob)) as probe:probe.verify()
                                png_structure_verified=True
                            except Exception:pass
                        byte_size_mb = len(blob) / (1024.0 * 1024.0)
                        real_size = get_image_dimensions(blob)
                        if real_size:
                            width_px, height_px = real_size
                    except Exception:
                        pass

            images_info.append({
                'index': idx,
                'xml_path':d.getroottree().getpath(d),
                'relationship_id':blip.get(qn('r:embed')) if blip is not None else None,
                'placement': placement,
                'native_display_width_mm':native_width,'native_display_height_mm':native_height,'native_display_area_mm2':native_area,'wrap_mode':wrap_mode,
                'ordinary_drawing_geometry_supported':extent is not None and len(d.findall('.//a:xfrm',NS))==1 and not d.find('.//a:xfrm',NS).attrib and len(d.findall('.//a:xfrm/a:ext',NS))==1 and d.find('.//a:xfrm/a:ext',NS).get('cx')==extent.get('cx') and d.find('.//a:xfrm/a:ext',NS).get('cy')==extent.get('cy') and d.find('.//a:srcRect',NS) is None,
                'extension': image_extension,
                'png_bit_depth':png_bit_depth,'png_color_type':png_color_type,
                'png_structure_verified':png_structure_verified,
                'blip_count':len(d.findall('.//a:blip',NS)),
                **zip_metadata,
                'width_px': width_px,
                'height_px': height_px,
                'display_width_px': display_width_px,
                'display_height_px': display_height_px,
                'byte_size_mb': byte_size_mb
            })
        return images_info

    def parse_structure_paragraphs(self):
        """Document-order text including table covers and SDT TOCs; no inferred page numbers."""
        main={p._p:i for i,p in enumerate(self.doc.paragraphs)}
        records=[]
        for element in self.doc.element.body.xpath('.//w:p'):
            if element.xpath('ancestor::w:del'):continue
            chunks=[]
            for node in element.xpath('.//w:t | .//w:tab | .//w:br'):
                # A drawing's nested textbox paragraph is a separate record, not duplicate text.
                nearest=next((a for a in node.iterancestors() if a.tag==qn('w:p')),None)
                if nearest is element:chunks.append(node.text or '' if node.tag==qn('w:t') else '\t' if node.tag==qn('w:tab') else '\n')
            records.append({'index':len(records),'text':''.join(chunks),'xml_path':element.getroottree().getpath(element),'main_paragraph_index':main.get(element),'inside_table':bool(element.xpath('ancestor::w:tbl')),'inside_content_control':bool(element.xpath('ancestor::w:sdt')),'inside_textbox':bool(element.xpath('ancestor::w:txbxContent'))})
        return records

    def parse_all(self) -> Dict[str, Any]:
        images=self.parse_drawings_and_images()
        labels=vml_textboxes(self.doc.element.body)
        labels_profile=(len(labels)==len(self.doc.element.body.findall('.//{urn:schemas-microsoft-com:vml}shape')) and all(l['geometry_supported'] for l in labels) and not bool(self.doc.element.xpath('.//*[local-name()="txbx"] | .//*[local-name()="group"]')))
        bounds_profile=(len(labels)==len(self.doc.element.body.findall('.//{urn:schemas-microsoft-com:vml}shape')) and all(l['bounds_geometry_supported'] for l in labels) and not bool(self.doc.element.xpath('.//*[local-name()="txbx"] | .//*[local-name()="group"]')))
        floating_profile=(1<=len(images)<=2 and not labels and all(im['placement'] in {'inline','floating'} and (im['placement']=='inline' or im['wrap_mode'] in {'wrapNone','wrapSquare'}) and im['extension']=='png' and im['png_structure_verified'] and im['png_bit_depth']==8 and im['png_color_type']==2 and im['blip_count']==1 and im['native_display_area_mm2'] is not None and im['ordinary_drawing_geometry_supported'] for im in images) and not bool(self.doc.element.xpath('.//w:pict | .//*[local-name()="AlternateContent"]')))
        bytes_profile=(1<=len(images)<=2 and all(im.get('placement')=='inline' and im.get('extension')=='png' and im.get('png_bit_depth')==8 and im.get('png_color_type')==2 and im.get('png_structure_verified') and im.get('blip_count')==1 and im.get('zip_local_flags') in {0,8} and im.get('zip_compression_method')==8 and not im.get('zip_uses_zip64',True) and not im.get('zip_entry_ambiguous',True) and (im.get('zip_local_uncompressed_size')==im.get('zip_central_uncompressed_size') if im.get('zip_local_flags')==0 else im.get('zip_local_uncompressed_size')==0) for im in images)
                       and not bool(self.doc.element.xpath('.//w:pict | .//w:txbxContent | .//w:object | .//*[local-name()="AlternateContent"]')))
        pixel_profile=(len(images)==1 and images[0]['placement']=='inline' and images[0]['extension']=='png' and images[0]['png_bit_depth']==8 and images[0]['png_color_type']==2 and images[0]['png_structure_verified'] and images[0]['blip_count']==1
                       and all(isinstance(images[0].get(k),(int,float)) and images[0][k]>0 for k in ['width_px','height_px'])
                       and images[0]['width_px']*images[0]['height_px']<2**31
                       and not bool(self.doc.element.xpath('.//w:pict | .//w:txbxContent | .//w:object | .//*[local-name()="AlternateContent"]')))
        return {
            'labels':labels,'native232_vml_label_profile_supported':labels_profile,'native232_vml_bounds_profile_supported':bounds_profile,'native232_floating_profile_supported':floating_profile,
            'native232_image_bytes_profile_supported':bytes_profile,
            'native232_single_inline_png_profile_supported':pixel_profile,
            'native_graduate_scope_supported': not bool(self.doc.element.xpath('.//w:tbl | .//w:drawing | .//w:pict | .//w:txbxContent | .//w:fldSimple | .//w:instrText | .//w:del | .//w:ins | .//w:altChunk | .//w:sdt')),
            'formatting_coverage': {'resolved': ['docDefaults', 'paragraph_basedOn', 'character_basedOn', 'direct'], 'unresolved': ['conditional_table_styles', 'numbering_properties', 'theme_font_and_color'], 'drawing_geometry':{'measured_vml_textboxes':sum(label['geometry_supported'] for label in labels),'unmeasured_vml_textboxes':sum(not label['geometry_supported'] for label in labels),'floating_png_profile_supported':floating_profile,'vml_label_profile_supported':labels_profile,'limits':'Simple explicit geometry only; general shapes, DrawingML textboxes, fit/crop/rotation and other units require separate proof.'}},
            'file_size_mb': self.metadata['file_size_mb'],
            'sections': self.metadata['sections'],
            'paragraphs': self.parse_paragraphs(),
            'structure_paragraphs': self.parse_structure_paragraphs(),
            'tables': self.parse_tables(),
            'images': images
        }
