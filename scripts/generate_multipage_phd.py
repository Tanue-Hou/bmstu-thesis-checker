"""Frozen multipage PDF and doctoral copy-repair counterfactual inputs."""
import hashlib,json,shutil,zipfile
from pathlib import Path
from io import BytesIO
import fitz
from PIL import Image
from docx import Document
from docx.shared import Pt,RGBColor,Mm
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'evidence/multipage-phd-repair';STAGE=Path('C:/Users/Tanue Hou/AppData/Local/TestVkrLab-20261014')
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def record(p,params,parent=None):
    shutil.copyfile(p,STAGE/p.name)
    r={'case':p.stem,'format':p.suffix[1:],'sha256':sha(p),'parameters':params}
    if parent:r.update(parent=str(parent.relative_to(ROOT)),parent_sha256=sha(parent))
    return r
def fonts(page):page.insert_font(fontname='TNR',fontfile='C:/Windows/Fonts/times.ttf')
def make_multi(name,fontsize=14,small=False,red=False,letter=False,largeimage=False):
    doc=fitz.open();chart=fitz.open(ROOT/'evidence/format-expansion/pdf_complex_clean.pdf');image=chart.extract_image(chart[0].get_images()[0][0])['image'];chart.close();small_image=image
    if largeimage:
        b=BytesIO();Image.new('RGB',(4001,4000),'white').save(b,format='PNG');image=b.getvalue()
    for pi in range(3):
        p=doc.new_page(width=(612 if letter and pi==2 else 595.2756),height=(792 if letter and pi==2 else 841.8898));fonts(p)
        for i,t in enumerate(['НАУЧНО-КВАЛИФИКАЦИОННАЯ РАБОТА','ОГЛАВЛЕНИЕ','ВВЕДЕНИЕ','Методы обработки данных и анализа.','Результаты исследования представлены.','ЗАКЛЮЧЕНИЕ','СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ']):p.insert_text((72,65+i*35),t,fontname='TNR',fontsize=14,color=(1,0,0) if red and pi==2 and i==3 else (0,0,0))
        for row in range(3):
            for col in range(2):
                x,y=72+col*190,330+row*27;p.draw_rect(fitz.Rect(x,y,x+190,y+27),color=(0,0,0),width=.5);p.insert_text((x+8,y+18),'Метод анализа данных',fontname='TNR',fontsize=7 if small and pi==1 else 14)
        p.insert_image(fitz.Rect(72,440,500,615),stream=image if pi==1 else small_image)
        if pi==1:
            for i in range(6):p.insert_text((72+i*70,735),'Да',fontname='TNR',fontsize=fontsize)
        p.insert_text((280,800),str(pi+1),fontname='TNR',fontsize=14)
    path=OUT/(name+'.pdf');doc.save(path,garbage=4,deflate=True);doc.close();return record(path,{'pages':3,'table_small_page':2 if small else None,'red_page':3 if red else None,'letter_page':3 if letter else None,'large_image_page':2 if largeimage else None,'probe_font_pt':fontsize})
def make_structure(name,drop_abstract=False,wrong_toc=False):
    original=ROOT/'evidence/graduate-structure/g_phd_left25.docx';texts=[p.text for p in Document(original).paragraphs];doc=fitz.open();page=None;y=0;font=fitz.Font(fontfile='C:/Windows/Fonts/times.ttf')
    def new():
        nonlocal page,y
        page=doc.new_page();fonts(page);y=60
    new()
    for text in texts:
        if drop_abstract and text=='РЕФЕРАТ':continue
        if wrong_toc and text=='ОГЛАВЛЕНИЕ':text='СОДЕРЖАНИЕ'
        if not text:continue
        if text in ('РЕФЕРАТ','ОГЛАВЛЕНИЕ','СОДЕРЖАНИЕ','ВВЕДЕНИЕ','1 ОСНОВНАЯ ЧАСТЬ','ЗАКЛЮЧЕНИЕ','СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ','ПРИЛОЖЕНИЕ А') and y>100:new()
        line=''
        for word in text.split():
            candidate=(line+' '+word).strip()
            if font.text_length(candidate,fontsize=14)>450 and line:
                if y>760:new()
                page.insert_text((71,y),line,fontname='TNR',fontsize=14);y+=20;line=word
            else:line=candidate
        if line:
            if y>760:new()
            page.insert_text((71,y),line,fontname='TNR',fontsize=14);y+=26
    p=OUT/(name+'.pdf');count=len(doc);doc.save(p,garbage=4,deflate=True);doc.close();return record(p,{'pages':count,'docx_text_flow_conversion':True,'drop_abstract_heading':drop_abstract,'wrong_toc_heading':wrong_toc},original)
def main():
    OUT.mkdir(exist_ok=True)
    if (OUT/'inputs.json').exists():raise SystemExit('Frozen inputs exist')
    rows=[]
    for name,kw in [('mp_clean',{}),('mp_small',{'small':True}),('mp_red',{'red':True}),('mp_letter',{'letter':True}),('mp_image_above',{'largeimage':True}),('mp_font18_5',{'fontsize':18.5}),('mp_font19',{'fontsize':19}),('mp_font24',{'fontsize':24}),('mp_font36',{'fontsize':36}),('mp_font72',{'fontsize':72})]:rows.append(make_multi(name,**kw))
    for name,kw in [('structure_clean',{}),('structure_no_abstract',{'drop_abstract':True}),('structure_wrong_toc',{'wrong_toc':True})]:rows.append(make_structure(name,**kw))
    parent=ROOT/'evidence/format-expansion/complex_clean.docx'
    for name,mode in [('phd_complex_control','clean'),('phd_small','small'),('phd_large','large'),('phd_red','red'),('phd_letter','letter'),('phd_combined','combined')]:
        d=Document(parent)
        # Insert long ordinary body text before the first table; preserve chart/table parts.
        if mode in ('small','large','red','combined'):
            probes=[('small',7),('large',18.5),('red',14)] if mode=='combined' else [(mode,7 if mode=='small' else 18.5 if mode=='large' else 14)]
            for kind,size in probes:
                p=d.add_paragraph('Методы обработки данных и анализа результатов исследования позволяют проверить форматирование текста.');d.tables[0]._tbl.addprevious(p._p)
                for run in p.runs:run.font.size=Pt(size);run.font.color.rgb=RGBColor(255,0,0) if kind=='red' else RGBColor(0,0,0)
        if mode in ('letter','combined'):d.sections[0].page_width=Mm(216);d.sections[0].page_height=Mm(279)
        path=OUT/(name+'.docx');d.save(path);rows.append(record(path,{'probe':mode,'tables_and_inline_chart':True},parent))
    (OUT/'inputs.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    for name in ['mp_clean','mp_font72','structure_clean']:
        with fitz.open(OUT/(name+'.pdf')) as d:
            for pi in range(len(d)):d[pi].get_pixmap(matrix=fitz.Matrix(.8,.8)).save(str(OUT/(name+f'.render-{pi+1}.png')))
    print(json.dumps({'inputs':len(rows),'pdf':sum(r['format']=='pdf' for r in rows)}))
if __name__=='__main__':main()
