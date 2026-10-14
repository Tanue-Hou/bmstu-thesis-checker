"""Reproducible native probes; generated content is synthetic, not a real thesis."""
import hashlib,json,shutil,zipfile
from pathlib import Path
from io import BytesIO
from docx import Document
from docx.shared import RGBColor
from docx.oxml import OxmlElement
from lxml import etree
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'evidence/graduate-structure'
STAGE=Path('C:/Users/Tanue Hou/AppData/Local/TestVkrLab-20261014')

def save(doc,name,pages=None):
    buf=BytesIO();doc.save(buf);buf.seek(0);p=OUT/(name+'.docx')
    with zipfile.ZipFile(buf) as src,zipfile.ZipFile(p,'w') as dst:
        for info in src.infolist():
            data=src.read(info.filename)
            if info.filename=='docProps/app.xml' and pages is not None:
                root=etree.fromstring(data);root.find('{*}Pages').text=str(pages);data=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
            dst.writestr(info,data)
    shutil.copyfile(p,STAGE/p.name)
    return {'case':name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'cached_pages':pages}

def base(): return Document(ROOT/'evidence/dynamic-master/repair_combined_after.docx')

def cover(doc,phd=False):
    end=next(p for p in doc.paragraphs if p.text=='РЕФЕРАТ')
    for p in list(doc.paragraphs):
        if p==end: break
        if p._p is end._p: break
        p._p.getparent().remove(p._p)
    lines=['Министерство науки и высшего образования Российской Федерации',
      'Федеральное государственное автономное образовательное учреждение высшего образования',
      'Московский государственный технический университет имени Н. Э. Баумана',
      '(национальный исследовательский университет) (МГТУ им. Н. Э. Баумана)',
      'ФАКУЛЬТЕТ Информатика и системы управления','КАФЕДРА Информационные системы',
      'НАУЧНО-КВАЛИФИКАЦИОННАЯ РАБОТА' if phd else 'ВЫПУСКНАЯ КВАЛИФИКАЦИОННАЯ РАБОТА МАГИСТРА',
      'НА ТЕМУ: Исследование методов обработки данных',
      'Аспирант: Иванов Иван Иванович' if phd else 'Студент: Иванов Иван Иванович',
      'Руководитель: Петров Петр Петрович','Нормоконтролер: Сидоров Сидор Сидорович','Москва 2026']
    for text in lines: end.insert_paragraph_before(text)

def complete(doc,phd=False):
    cover(doc,phd)
    ps=doc.paragraphs;start=next(i for i,p in enumerate(ps) if p.text=='РЕФЕРАТ');end=next(p for p in ps if p.text=='СОДЕРЖАНИЕ')
    for p in ps[start+1:]:
        if p._p is end._p:break
        p._p.getparent().remove(p._p)
    for _ in range(3):end.insert_paragraph_before('В работе рассматриваются методы обработки данных и анализа результатов исследования. '*3)
    # Add actual numbered sections and matching TOC rows to cross the native 11-entry threshold.
    intro=next(p for p in doc.paragraphs if p.text=='ВВЕДЕНИЕ')
    toc_start=next(i for i,p in enumerate(doc.paragraphs) if p.text=='СОДЕРЖАНИЕ')
    for p in list(doc.paragraphs)[toc_start+1:]:
        if p._p is intro._p:break
        p._p.getparent().remove(p._p)
    rows=['ВВЕДЕНИЕ','1 ОСНОВНАЯ ЧАСТЬ']+[f'1.{i} Методы исследования {i}' for i in range(1,8)]+['ЗАКЛЮЧЕНИЕ','СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ','ПРИЛОЖЕНИЕ А']
    for i,text in enumerate(rows,4):intro.insert_paragraph_before(text+' .......... '+str(i))
    conclusion=next(p for p in doc.paragraphs if p.text=='ЗАКЛЮЧЕНИЕ')
    for i in range(1,8):
        conclusion.insert_paragraph_before(f'1.{i} Методы исследования {i}','Heading 1')
        for _ in range(3):conclusion.insert_paragraph_before('Методы исследования позволяют получить результаты обработки данных и выполнить анализ. '*3)
    doc.add_paragraph('ПРИЛОЖЕНИЕ А','Heading 1');doc.add_paragraph('Дополнительные результаты исследования. '*10)
    for p in doc.paragraphs:
        for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
    return doc

if __name__=='__main__':
    OUT.mkdir(parents=True,exist_ok=True)
    if (OUT/'inputs.json').exists():raise SystemExit('Frozen input manifest exists')
    cases=[save(base(),'g_baseline')]
    d=base();cover(d);cases.append(save(d,'g_cover'))
    d=complete(base());cases.append(save(d,'g_structure'))
    # Large body with unchanged explicit breaks; test cached page metadata independently.
    anchor=next(p for p in d.paragraphs if p.text=='ЗАКЛЮЧЕНИЕ')
    for _ in range(75):anchor.insert_paragraph_before('Исследование методов обработки данных и анализ результатов. '*12)
    cases.extend([save(d,'g_pages47',47),save(d,'g_pages48',48)])
    dp=complete(base(),True);cases.append(save(dp,'g_phd_structure'))
    (OUT/'inputs.json').write_text(json.dumps(cases,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'cases':len(cases)}))
