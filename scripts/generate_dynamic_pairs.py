"""Controlled thesis-shaped inputs; one mutated property per case."""
import hashlib
import json
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'evidence' / 'dynamic-pairs'

def create(name, spacing=1.5, indent=1.25, font=14, a4=True):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = (Mm(210), Mm(297)) if a4 else (Mm(216), Mm(279))
    sec.left_margin, sec.right_margin = Cm(3), Cm(1)
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.header_distance = sec.footer_distance = Cm(1.25)
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(14)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for s in ('Heading 1', 'Title'):
        doc.styles[s].font.name = 'Times New Roman'
        doc.styles[s].font.size = Pt(14)
        doc.styles[s].font.bold = True
    doc.add_paragraph('МГТУ им. Н. Э. Баумана', 'Title')
    doc.add_paragraph('ВЫПУСКНАЯ КВАЛИФИКАЦИОННАЯ РАБОТА БАКАЛАВРА', 'Title')
    doc.add_paragraph('Исследование методов обработки данных')
    doc.add_paragraph('Студент: Иванов Иван Иванович')
    doc.add_page_break()
    doc.add_paragraph('РЕФЕРАТ', 'Heading 1')
    doc.add_paragraph('Работа посвящена исследованию методов обработки данных. Результаты представлены в основной части работы. ' * 5)
    doc.add_page_break()
    doc.add_paragraph('СОДЕРЖАНИЕ', 'Heading 1')
    for line in ('ВВЕДЕНИЕ ........................................ 4', '1 ОСНОВНАЯ ЧАСТЬ ............................... 5',
                 'ЗАКЛЮЧЕНИЕ ..................................... 6', 'СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ ............... 7'):
        doc.add_paragraph(line)
    for heading in ('ВВЕДЕНИЕ', '1 ОСНОВНАЯ ЧАСТЬ', 'ЗАКЛЮЧЕНИЕ'):
        doc.add_page_break()
        doc.add_paragraph(heading, 'Heading 1')
        for _ in range(6):
            p = doc.add_paragraph('В работе рассматриваются методы обработки данных и анализа результатов исследования. ' * 8 + 'Методы описаны в источнике [1].')
            p.paragraph_format.line_spacing = spacing
            p.paragraph_format.first_line_indent = Cm(indent)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(font)
    doc.add_page_break()
    doc.add_paragraph('СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ', 'Heading 1')
    doc.add_paragraph('1. Иванов И. И. Методы обработки данных. Москва, 2024. 200 с.')
    path = OUT / (name + '.docx')
    doc.save(path)
    return {'case': name, 'path': str(path), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
            'parameters': {'spacing': spacing, 'indent_cm': indent, 'font_pt': font, 'a4': a4}}

if __name__ == '__main__':
    import argparse
    from io import BytesIO
    from PIL import Image
    from docx.shared import RGBColor
    parser=argparse.ArgumentParser()
    parser.add_argument('--output-dir',type=Path,default=OUT)
    args=parser.parse_args()
    OUT=args.output_dir.resolve()
    if (OUT/'inputs.json').exists():
        raise SystemExit('Refusing to overwrite frozen inputs. Choose a new --output-dir.')
    OUT.mkdir(parents=True, exist_ok=True)
    cases = [create('baseline'), create('spacing_1', spacing=1), create('indent_0', indent=0),
             create('font_7', font=7), create('page_letter', a4=False)]
    cases.extend(create(name,font=value) for name,value in [('font_7_5',7.5),('font_8',8),('font_18',18),('font_18_5',18.5)])
    for name,width,height in [('image_16m',4000,4000),('image_above_16m',4001,4000)]:
        doc=Document(OUT/'baseline.docx')
        stream=BytesIO();Image.new('RGB',(width,height),'white').save(stream,format='PNG');stream.seek(0)
        doc.add_picture(stream,width=Cm(10));path=OUT/(name+'.docx');doc.save(path)
        cases.append({'case':name,'path':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
                      'parameters':{'width_px':width,'height_px':height}})
    doc=Document(OUT/'baseline.docx')
    for style in doc.styles:
        if style.type==1:style.font.color.rgb=RGBColor(0,0,0)
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:run.font.color.rgb=RGBColor(0,0,0)
    path=OUT/'color_black.docx';doc.save(path)
    cases.append({'case':'color_black','path':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
                  'parameters':{'text_color':'000000'}})
    (OUT / 'inputs.json').write_text(json.dumps(cases, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'cases': len(cases), 'output': str(OUT)}))
