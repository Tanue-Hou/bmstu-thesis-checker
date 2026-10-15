import argparse
import hashlib
import json
import os
import sys
from pathlib import Path


SERVICE_ROOT = Path(__file__).resolve().parent / "testvkr-service"
sys.path.insert(0, str(SERVICE_ROOT))

from checkers.dictionary import DictionaryChecker
from checkers.entry_document import EntryDocumentChecker
from checkers.pictures import PicturesChecker
from checkers.structure import StructureChecker
from checkers.word_format import WordFormatChecker
from parsers.docx_parser import DocxParser
from parsers.pdf_parser import PdfParser
from config import settings
from school_templates import load_template,check_template
from issue_locations import locate_finding
from native_builds import require_build


SEVERITY_ORDER = {"critical": 0, "major": 1, "warning": 2, "info": 3}
FIX_PRIORITY_ORDER = {"must_fix": 0, "should_fix": 1, "manual_review": 2, "optional": 3}
SEVERITY_LABELS_ZH = {
    "critical": "严重",
    "major": "较严重",
    "warning": "警告",
    "info": "提示",
}
FIX_PRIORITY_LABELS_ZH = {
    "must_fix": "必改",
    "should_fix": "建议改",
    "manual_review": "需要人工确认",
    "optional": "可忽略",
}
MODULE_LABELS_ZH = {
    "entry": "入口/文件",
    "structure": "结构",
    "pictures": "图片/图形",
    "word_format": "Word/PDF 格式",
    "dictionary": "拼写/词典",
}


def load_rule_registry():
    rules_path = SERVICE_ROOT / "rules" / "rules_verified.json"
    if not rules_path.exists():
        return {}
    with rules_path.open("r", encoding="utf-8-sig") as f:
        data = json.load(f)
    registry = {rule["rule_id"]: rule for rule in data.get("rules", [])}
    # The runtime policy is authoritative; registry descriptions must not drift
    # when users supply environment overrides through BaseSettings.
    bindings = {
        "Word.LineSpacingInvalid": "LINE_SPACING_BODY",
        "Word.FontSizeTooSmall": "FONT_SIZE_MIN_PT",
        "Word.FontSizeTooLarge": "FONT_SIZE_MAX_PT",
        "Word.FontScaleTooSmall": "FONT_SCALE_MIN_PCT",
        "Word.FontScaleTooLarge": "FONT_SCALE_MAX_PCT",
        "Word.TableRowHeightTooSmall": "TABLE_ROW_HEIGHT_MIN_MM",
        "Word.HeaderFooterDistanceInvalid": "HEADER_FOOTER_DISTANCE_CM",
        "Word.FirstLineIndentInvalid": "FIRST_LINE_INDENT_CM",
        "Word.LeftMarginInvalid": "MARGIN_LEFT_MM",
        "Word.RightMarginInvalid": "MARGIN_RIGHT_MM",
        "Word.TopMarginInvalid": "MARGIN_TOP_MM",
        "Word.BottomMarginInvalid": "MARGIN_BOTTOM_MM",
        "Pic.FileSizeTooLarge": "IMAGE_FILE_SIZE_MAX_MB",
        "Pic.PixelSizeTooLarge": "IMAGE_PIXEL_AREA_MAX",
        "Struct.ThesisSizeTooSmall": "THESIS_SIZE_MIN_BACHELOR",
    }
    for rule_id, key in bindings.items():
        if rule_id in registry:
            registry[rule_id]["threshold"] = getattr(settings, key)
            registry[rule_id]["policy_setting"] = key
    if "Pic.PixelSizeTooLarge" in registry:
        registry["Pic.PixelSizeTooLarge"]["enabled_by_default"] = True
    for rule in registry.values():
        rule["comparison_policy"] = {
            "status": "local_policy_not_original_program_proof",
            "line_spacing_tolerance": settings.LINE_SPACING_TOLERANCE,
            "table_line_spacing": settings.LINE_SPACING_TABLE,
            "font_violation_min_chars_exclusive": settings.FONT_VIOLATION_MIN_CHARS,
            "page_size_tolerance_mm": settings.PAGE_SIZE_TOLERANCE_MM,
            "margin_tolerance_mm": settings.MARGIN_TOLERANCE_MM,
            "header_footer_tolerance_cm": settings.HEADER_FOOTER_TOLERANCE_CM,
            "first_line_indent_tolerance_cm": settings.FIRST_LINE_INDENT_TOLERANCE_CM,
        }
    return registry


def load_rule_guidance():
    guidance_path = SERVICE_ROOT / "rules" / "rule_guidance_zh.json"
    if not guidance_path.exists():
        return {"default": {}}
    with guidance_path.open("r", encoding="utf-8-sig") as f:
        return json.load(f)


def parse_document(path: Path):
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return "docx", DocxParser(str(path)).parse_all()
    if suffix == ".pdf":
        return "pdf", PdfParser(str(path)).parse_all()
    raise ValueError("Only .docx and .pdf files are supported.")


def describe_location(location):
    if not location:
        return "未提供精确位置；请结合模块和阈值信息人工定位。"

    parts = []
    if "section" in location:
        parts.append(f"section {location['section']}")
    if "page" in location:
        parts.append(f"PDF/页面 {location['page']}")
    if "pages" in location:
        parts.append("PDF/页面 "+", ".join(str(p) for p in location["pages"]))
    if "paragraph" in location:
        parts.append(f"段落 {location['paragraph']}")
    if "paragraphs" in location:
        values = location["paragraphs"]
        if isinstance(values, list):
            shown = ", ".join(str(v) for v in values[:20])
            suffix = "..." if len(values) > 20 else ""
            parts.append(f"段落 {shown}{suffix}")
    if "table" in location:
        parts.append(f"表格 {location['table']}")
    if "tables" in location:
        values = location["tables"]
        if isinstance(values, list):
            shown = ", ".join(str(v) for v in values[:20])
            suffix = "..." if len(values) > 20 else ""
            parts.append(f"表格 {shown}{suffix}")
    if "row" in location:
        parts.append(f"行 {location['row']}")
    if "table_rows" in location:
        rows = location["table_rows"]
        if isinstance(rows, list):
            shown = ", ".join(f"表{r.get('table')}行{r.get('row')}" for r in rows[:20])
            suffix = "..." if len(rows) > 20 else ""
            parts.append(f"表格行 {shown}{suffix}")
    if "image" in location:
        parts.append(f"图片 {location['image']}")
    if "images" in location:
        values = location["images"]
        if isinstance(values, list):
            shown = ", ".join(str(v) for v in values[:20])
            suffix = "..." if len(values) > 20 else ""
            parts.append(f"图片 {shown}{suffix}")
    if "bbox" in location:
        parts.append(f"bbox {location['bbox']}")
    if "text" in location:
        parts.append(f"文本片段：{location['text']}")

    return "；".join(parts) if parts else json.dumps(location, ensure_ascii=False)


def location_group_key(location):
    if not location:
        return "未定位"
    if "page" in location:
        return f"页面 {location['page']}"
    if "pages" in location:
        return "页面 "+", ".join(str(p) for p in location["pages"])
    if "section" in location:
        return f"section {location['section']}"
    if "paragraph" in location or "paragraphs" in location:
        return "段落"
    if "table" in location or "tables" in location or "table_rows" in location:
        return "表格"
    if "image" in location or "images" in location:
        return "图片"
    return "其他位置"


def verification_context(rule,file_type,work_type):
    if 'verification_by_context' in rule:
        return rule['verification_by_context'].get(file_type+':'+work_type) if file_type and work_type else None
    if file_type not in rule.get('verified_formats',[]) or work_type not in rule.get('verified_work_types',[]):return None
    return rule.get('verification_by_work_type',{}).get(work_type,rule)

def context_evidence_valid(context):
    if not context:return False
    refs=context.get('evidence_refs',[])
    if not refs:return False
    try:
        root=Path(__file__).resolve().parents[1]
        if hashlib.sha256((root/refs[0]).read_bytes()).hexdigest()!=context.get('verification_evidence_sha256'):return False
        return all(hashlib.sha256((root/ref).read_bytes()).hexdigest()==context.get('evidence_hashes',{}).get(ref,context.get('verification_evidence_sha256')) for ref in refs)
    except OSError:return False


def finding_to_dict(finding, rule_registry, rule_guidance, file_type=None, work_type=None, doc_data=None):
    rule = rule_registry.get(finding.code, {})
    selected_context=verification_context(rule,file_type,work_type)
    context = selected_context or {}
    verified = (rule.get('verification_status') == 'verified_against_original'
                and (doc_data or {}).get('native_build',230)==230
                and selected_context is not None
                and (doc_data or {}).get("document_kind")!="dissertation"
                and context_evidence_valid(selected_context)
                and (file_type!='pdf' or bool((doc_data or {}).get('pdf_structure_profile_supported' if finding.code=='Struct.RequiredHeadingsMissing' else 'pdf_verification_profile_supported')))
                and (finding.code!='Struct.RequiredHeadingsMissing' or bool((finding.threshold or {}).get('native_predicate_supported')))
                and (finding.code not in {'Word.LeftMarginInvalid','Word.RightMarginInvalid','Word.TopMarginInvalid','Word.BottomMarginInvalid'} or work_type!='phd' or bool((finding.threshold or {}).get('native_margin_probe_supported')))
                and (not finding.code.startswith('Graduate.') or (bool((finding.threshold or {}).get('native_scope_supported')) and bool((finding.threshold or {}).get('native_predicate_supported'))))
                and file_type in rule.get('verified_formats', [])
                and work_type in rule.get('verified_work_types', [])
                and all(getattr(settings, key, None) == value
                        for key, value in context.get('verified_policy',rule.get('verified_policy', {})).items()))
    verification_status = ('verified_against_original' if verified else
                           'pending_original_program_comparison')
    guidance = {**rule_guidance.get("default", {}), **rule_guidance.get(finding.code, {})}
    location = finding.location
    return {
        "rule_id": finding.code,
        "module": finding.module.value,
        "severity": finding.severity.value,
        "message": finding.message,
        "russian_message": finding.message,
        "title_zh": guidance.get("title_zh"),
        "explanation_zh": guidance.get("explanation_zh"),
        "trigger_reason_zh": guidance.get("trigger_reason_zh"),
        "fix_suggestion_zh": guidance.get("fix_suggestion_zh"),
        "word_path_zh": guidance.get("word_path_zh"),
        "fix_priority": (guidance.get("fix_priority", "manual_review")
                         if verified
                         else "manual_review"),
        "suggested_fix_priority": guidance.get("fix_priority", "manual_review"),
        "affects_antiplagiarism": bool(guidance.get("affects_antiplagiarism")),
        "confidence": rule.get("confidence", "low") if verified else "low",
        "verification_status": verification_status,
        "verification_evidence_intact": context_evidence_valid(selected_context),
        "verification_scope": context.get('scope', rule.get("verification_scope", {})) if selected_context is not None else {},
        "threshold_origin": rule.get("threshold_origin", "local_heuristic"),
        "evidence_refs": context.get('evidence_refs', rule.get("evidence_refs", [])) if selected_context is not None else [],
        "static_evidence_refs": rule.get("static_evidence_refs", []),
        "comparison_policy": {**rule.get("comparison_policy", {}), **({"margin_tolerance_mm": settings.MARGIN_TOLERANCE_PHD_MM} if work_type=="phd" and finding.code in {"Word.LeftMarginInvalid","Word.RightMarginInvalid","Word.TopMarginInvalid","Word.BottomMarginInvalid"} else {})},
        "threshold": finding.threshold,
        "location": location,
        "location_summary_zh": describe_location(location),
        "skip_switch": finding.skip_switch,
        "source_va": rule.get("source_va"),
        "comparator": rule.get("comparator"),
        "sample_trigger": rule.get("sample_trigger", []),
    }


def build_groups(findings):
    groups = {
        "by_severity": {},
        "by_module": {},
        "by_fix_priority": {},
        "by_location": {},
        "antiplagiarism_related": [],
    }
    for item in findings:
        groups["by_severity"].setdefault(item["severity"], []).append(item["rule_id"])
        groups["by_module"].setdefault(item["module"], []).append(item["rule_id"])
        groups["by_fix_priority"].setdefault(item["fix_priority"], []).append(item["rule_id"])
        groups["by_location"].setdefault(location_group_key(item.get("location")), []).append(item["rule_id"])
        if item["affects_antiplagiarism"]:
            groups["antiplagiarism_related"].append(item["rule_id"])
    return groups


def build_fix_checklist(findings):
    checklist = {key: [] for key in FIX_PRIORITY_LABELS_ZH}
    for item in findings:
        key = item.get("fix_priority") or "manual_review"
        checklist.setdefault(key, []).append({
            "rule_id": item["rule_id"],
            "label_zh": FIX_PRIORITY_LABELS_ZH.get(key, key),
            "title_zh": item.get("title_zh"),
            "severity": item.get("severity"),
            "module": item.get("module"),
            "location_summary_zh": item.get("location_summary_zh"),
            "fix_suggestion_zh": item.get("fix_suggestion_zh"),
            "word_path_zh": item.get("word_path_zh"),
            "affects_antiplagiarism": item.get("affects_antiplagiarism"),
        })
    return checklist


def build_review_summary(findings):
    must_fix = [f for f in findings if f.get("fix_priority") == "must_fix"]
    antiplagiarism = [f for f in findings if f.get("affects_antiplagiarism")]
    critical_or_major = [f for f in findings if f.get("severity") in {"critical", "major"}
                         and f.get("verification_status") == "verified_against_original"]
    if must_fix or critical_or_major:
        status = "不建议直接提交，需要先处理必改项和较高严重度问题。"
    elif findings:
        status = "发现本地策略提示，请先复核适用要求；这些提示尚不代表原版 TestVkr 的检查结论。"
    else:
        status = "未发现当前规则可识别的问题。"
    return {
        "status_zh": status,
        "must_fix_count": len(must_fix),
        "should_fix_count": sum(1 for f in findings if f.get("fix_priority") == "should_fix"),
        "manual_review_count": sum(1 for f in findings if f.get("fix_priority") == "manual_review"),
        "optional_count": sum(1 for f in findings if f.get("fix_priority") == "optional"),
        "antiplagiarism_related_count": len(antiplagiarism),
        "critical_or_major_count": len(critical_or_major),
    }


def sorted_findings(findings):
    return sorted(
        findings,
        key=lambda item: (
            SEVERITY_ORDER.get(item.get("severity"), 99),
            FIX_PRIORITY_ORDER.get(item.get("fix_priority"), 99),
            item.get("module", ""),
            item.get("rule_id", ""),
        ),
    )


def check_file(path: Path, args):
    rule_registry = load_rule_registry()
    rule_guidance = load_rule_guidance()
    file_type, doc_data = parse_document(path)
    template=load_template(getattr(args,"template_id",None),getattr(args,"template_path",None),args.work_type)
    kind=getattr(args,"document_kind","auto")
    if kind not in {"auto","vkr","nkr","dissertation"}:raise ValueError("Invalid document_kind")
    cover=" ".join(s.get("text","") for s in doc_data.get("spans",[]) if s.get("page_index")==0) if file_type=="pdf" else " ".join(p.get("text","") for p in doc_data.get("structure_paragraphs",doc_data.get("paragraphs",[]))[:35])
    import re
    detected_dissertation=bool(re.search(r"(?:кандидата|доктора)\s+.{0,45}?наук",cover.lower()))
    if kind=="auto":kind="dissertation" if detected_dissertation else "nkr" if args.work_type=="phd" else "vkr"
    if template and template["document_kind"]!=kind:raise ValueError("Template document kind does not match detected/selected document")
    doc_data["document_kind"]=kind
    native_build=getattr(args,"native_build",230)
    identity=require_build(native_build)
    doc_data["native_build"]=native_build

    findings = []
    checkers = [EntryDocumentChecker(doc_data, file_type, args.work_type)]

    entry_findings = checkers[0].check()
    findings.extend(entry_findings)
    stop = any(f.code in {"Entry.DocumentEmpty", "Entry.DocumentProtected"} for f in entry_findings)

    if not stop:
        if not args.skip_structure:
            findings.extend(StructureChecker(doc_data, file_type, args.work_type).check())
        if not args.skip_pictures:
            findings.extend(PicturesChecker(doc_data, file_type, args.work_type).check())
        if not args.skip_word_format:
            findings.extend(WordFormatChecker(doc_data, file_type, args.work_type).check())
        if not args.skip_dictionary:
            findings.extend(DictionaryChecker(doc_data, file_type, args.work_type).check())

    template_findings,template_coverage=check_template(doc_data,file_type,template) if not stop else ([],{})
    findings.extend(template_findings)
    finding_dicts = [finding_to_dict(f, rule_registry, rule_guidance, file_type, args.work_type,doc_data) for f in findings]
    for f in finding_dicts:
        f["precise_locations"]=locate_finding(f,doc_data,file_type)
        if f["rule_id"].startswith("Template."):
            f.update(verification_status="school_template_reference_not_native_proof",threshold_origin="selected_template",confidence="medium",template_id=template["id"],template_sources=template["sources"],applicability_confirmed=template["applicability_confirmed"],explanation_zh="与所选学校参考配置不同，需核对院系和年份。",fix_suggestion_zh="先确认该参考适用性，再在源文件中处理。",title_zh=f["message"])
    result = {
        "document_kind":kind,
        "native_reference":{"selected_build":native_build,"executable_sha256":identity["sha256"],"historical_rule_evidence_build":230,"transfer_validated":native_build==230,"limits":"Real-case family agreement is not controlled threshold validation"},
        "template":({k:v for k,v in template.items() if k!="targets"} if template else None),
        "template_coverage":template_coverage,
        "source_sha256":hashlib.sha256(Path(path).read_bytes()).hexdigest(),
        "location_contract":"PDF physical page+bbox; DOCX XML/paragraph/run/table coordinates, no inferred rendered page",
        "work_type_policy": ({"native_mode":None,"document_kind":"dissertation","manual_review":["Candidate/doctor of sciences dissertation: NKR native evidence and repair policy do not apply; use a confirmed council profile"]} if kind=="dissertation" else {"native_mode": "НКР аспиранта", "toc_heading": "ОГЛАВЛЕНИЕ", "docx_thesis_size_floor": "No numeric floor selected in the recovered build-230 branch; 100000-character heuristic is not applied.", "manual_review": ["Doctoral DOCX margins left/right/top/bottom 25/10/20/20 mm, tolerance 0.2 mm; finite single-section native probes", "Not a general dissertation standard or original-program certification"]} if args.work_type == 'phd' else {"native_mode": "ВКР магистра", "docx_size_measure": "48 equivalent pages; 900 trimmed UTF-16 characters per page; appendix A TOC page also checked", "manual_review": ["Use actual school template and real pagination; do not fabricate cover data, TOC rows or text"]} if args.work_type == 'master' else {}),
        "file": str(path),
        "file_type": file_type,
        "work_type": args.work_type,
        "verification_status": "local_policy_check_not_original_program_certification",
        "formatting_coverage": doc_data.get("formatting_coverage", {}),
        "native_comparison_limits": ({"format_profile_supported":doc_data.get("pdf_verification_profile_supported",False),"structure_profile_supported":doc_data.get("pdf_structure_profile_supported",False),"unimplemented_native_checks":["PDF text-derived margins and text in page margins","PDF page-number and header/footer placement","TOC correspondence and entry counts","cover/title-page predicates and bibliography completeness"],"upper_font_limit":"No native warning in tested 18.5/19/24/36/72 pt PDF probes; no universal upper-limit claim"} if file_type=="pdf" else {}),
        "summary": {
            "total": len(findings),
            "critical": sum(1 for f in findings if f.severity.value == "critical"),
            "major": sum(1 for f in findings if f.severity.value == "major"),
            "warning": sum(1 for f in findings if f.severity.value == "warning"),
            "info": sum(1 for f in findings if f.severity.value == "info"),
        },
        "review_summary": {**build_review_summary(finding_dicts),"native_verified_findings":sum(f["verification_status"]=="verified_against_original" for f in finding_dicts),"template_reference_findings":sum(f["rule_id"].startswith("Template.") for f in finding_dicts)},
        "groups": build_groups(finding_dicts),
        "fix_checklist": build_fix_checklist(finding_dicts),
        "findings": sorted_findings(finding_dicts),
    }
    return result


def render_markdown(result):
    lines = []
    def add(text=""):
        lines.append(text)

    add(f"# BMSTU 论文预审报告: {Path(result['file']).name}")
    add()
    add("本报告依据本地规则进行预审。未完成原程序对照的项目仅供复核，不作为已确认违规。")
    add()
    summary = result["summary"]
    review = result.get("review_summary", {})
    add(
        f"问题总数：{summary['total']} "
        f"（critical {summary['critical']}，major {summary['major']}，"
        f"warning {summary['warning']}，info {summary['info']}）"
    )
    add()
    if result.get("native_comparison_limits"):
        add("PDF 原程序检查覆盖仍不完整：边距、页眉页脚、页码位置和目录对应关系需要人工复核；零条本地问题不表示原程序通过。")
        add()
    if result.get("template"):
        add("学校参考配置："+result["template"]["id"]+"；当前适用性："+result["template_coverage"]["status"])
        add()
    add("## 审阅结论")
    add(f"- 结论：{review.get('status_zh', '')}")
    add(f"- 必改项：{review.get('must_fix_count', 0)}")
    add(f"- 建议改：{review.get('should_fix_count', 0)}")
    add(f"- 需要人工确认：{review.get('manual_review_count', 0)}")
    add(f"- 影响查重/反抄袭适配：{review.get('antiplagiarism_related_count', 0)}")
    add()

    add("## 分组汇总")
    for title, group_name, label_map in [
        ("按严重程度", "by_severity", SEVERITY_LABELS_ZH),
        ("按模块", "by_module", MODULE_LABELS_ZH),
        ("按位置", "by_location", {}),
    ]:
        add(f"### {title}")
        group = result.get("groups", {}).get(group_name, {})
        if not group:
            add("- 无")
        for key, rule_ids in group.items():
            add(f"- {label_map.get(key, key)}：{len(rule_ids)} 项")
        add()

    add("## 修复清单")
    for key, title in FIX_PRIORITY_LABELS_ZH.items():
        items = result.get("fix_checklist", {}).get(key, [])
        add(f"### {title}（{len(items)}）")
        if not items:
            add("- 无")
        for item in items:
            antiplagiarism = "；影响查重" if item.get("affects_antiplagiarism") else ""
            add(
                f"- `{item['rule_id']}` {item.get('title_zh') or ''}："
                f"{item.get('fix_suggestion_zh') or ''} "
                f"位置：{item.get('location_summary_zh') or '未提供'}{antiplagiarism}"
            )
        add()

    add("## 问题详情")
    for item in result["findings"]:
        add(f"### {item['severity'].upper()} - {item['rule_id']}")
        add(f"- 模块：{MODULE_LABELS_ZH.get(item['module'], item['module'])}")
        add(f"- 中文问题：{item.get('title_zh') or ''}")
        add(f"- 中文解释：{item.get('explanation_zh') or ''}")
        add(f"- 触发原因：{item.get('trigger_reason_zh') or ''}")
        add(f"- 修改建议：{item.get('fix_suggestion_zh') or ''}")
        add(f"- Word 操作路径：{item.get('word_path_zh') or ''}")
        add(f"- 位置：{item.get('location_summary_zh') or ''}")
        precise=item.get("precise_locations",{})
        add("- 精确匹配位置数："+str(precise.get("occurrence_count",0))+"；完整坐标见JSON，DOCX页码不由XML推测。")
        for occurrence in precise.get("occurrences",[])[:5]:
            add("  - `"+json.dumps(occurrence,ensure_ascii=False)+"`")
        add(f"- 是否影响查重/反抄袭：{'是' if item.get('affects_antiplagiarism') else '否'}")
        add(f"- 置信度：{item['confidence']}")
        add(f"- 验证状态：{item.get('verification_status', 'unverified')}")
        add(f"- 阈值来源：{item.get('threshold_origin', 'unknown')}")
        add(f"- 俄文原提示：{item['russian_message']}")
        if item.get("threshold") is not None:
            add(f"- 阈值：`{json.dumps(item['threshold'], ensure_ascii=False)}`")
        if item.get("location") is not None:
            add(f"- 原始位置：`{json.dumps(item['location'], ensure_ascii=False)}`")
        if item.get("source_va"):
            add(f"- 来源：`{item['source_va']}`")
        add()
    return "\n".join(lines)


def print_markdown(result):
    print(render_markdown(result))


def write_docx_report(result, out_path):
    from docx import Document
    from docx.shared import Pt

    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(10)

    document.add_heading(f"BMSTU 论文预审报告：{Path(result['file']).name}", level=1)
    document.add_paragraph("本报告依据本地规则进行预审。未完成原程序对照的项目仅供复核，不作为已确认违规。")
    summary = result["summary"]
    review = result.get("review_summary", {})
    document.add_paragraph(
        f"问题总数：{summary['total']}；严重：{summary['critical']}；"
        f"较严重：{summary['major']}；警告：{summary['warning']}；提示：{summary['info']}。"
    )
    document.add_heading("审阅结论", level=2)
    document.add_paragraph(review.get("status_zh", ""))
    document.add_paragraph(
        f"必改 {review.get('must_fix_count', 0)} 项；建议改 {review.get('should_fix_count', 0)} 项；"
        f"人工确认 {review.get('manual_review_count', 0)} 项；"
        f"影响查重/反抄袭适配 {review.get('antiplagiarism_related_count', 0)} 项。"
    )

    document.add_heading("修复清单", level=2)
    for key, title in FIX_PRIORITY_LABELS_ZH.items():
        items = result.get("fix_checklist", {}).get(key, [])
        document.add_heading(f"{title}（{len(items)}）", level=3)
        if not items:
            document.add_paragraph("无")
            continue
        table = document.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        headers = ["规则", "问题", "位置", "建议", "Word 操作路径"]
        for idx, header in enumerate(headers):
            table.rows[0].cells[idx].text = header
        for item in items:
            row = table.add_row().cells
            row[0].text = item.get("rule_id", "")
            row[1].text = item.get("title_zh") or ""
            row[2].text = item.get("location_summary_zh") or ""
            row[3].text = item.get("fix_suggestion_zh") or ""
            row[4].text = item.get("word_path_zh") or ""

    document.add_heading("问题详情", level=2)
    for item in result["findings"]:
        document.add_heading(f"{item['severity'].upper()} - {item['rule_id']}", level=3)
        details = [
            ("模块", MODULE_LABELS_ZH.get(item["module"], item["module"])),
            ("中文问题", item.get("title_zh") or ""),
            ("中文解释", item.get("explanation_zh") or ""),
            ("触发原因", item.get("trigger_reason_zh") or ""),
            ("修改建议", item.get("fix_suggestion_zh") or ""),
            ("Word 操作路径", item.get("word_path_zh") or ""),
            ("位置", item.get("location_summary_zh") or ""),
            ("是否影响查重/反抄袭", "是" if item.get("affects_antiplagiarism") else "否"),
            ("置信度", item.get("confidence") or ""),
            ("验证状态", item.get("verification_status") or "unverified"),
            ("阈值来源", item.get("threshold_origin") or "unknown"),
            ("俄文检查消息", item.get("russian_message") or ""),
        ]
        if item.get("threshold") is not None:
            details.append(("阈值", json.dumps(item["threshold"], ensure_ascii=False)))
        if item.get("source_va"):
            details.append(("来源", item["source_va"]))
        table = document.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for key, value in details:
            row = table.add_row().cells
            row[0].text = key
            row[1].text = str(value)

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out_path))
    return out_path


def write_report_bundle(result, report_dir):
    report_dir = Path(report_dir)
    report_dir.mkdir(parents=True, exist_ok=True)
    base = Path(result["file"]).stem
    paths = {
        "json": report_dir / f"{base}_report.json",
        "markdown": report_dir / f"{base}_report.md",
        "docx": report_dir / f"{base}_report.docx",
    }
    paths["json"].write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    paths["markdown"].write_text(render_markdown(result), encoding="utf-8")
    write_docx_report(result, paths["docx"])
    return {key: str(value) for key, value in paths.items()}


def main():
    parser = argparse.ArgumentParser(description="Check a DOCX/PDF thesis with the BMSTU rules engine.")
    parser.add_argument("file", help="Path to a .docx or .pdf file.")
    parser.add_argument("--work-type", default="bachelor", choices=["bachelor", "specialist", "master", "phd"])
    parser.add_argument("--native-build",type=int,choices=[230,232],default=232)
    parser.add_argument("--template-id")
    parser.add_argument("--template-path")
    parser.add_argument("--document-kind",choices=["auto","vkr","nkr","dissertation"],default="auto")
    parser.add_argument("--skip-structure", action="store_true")
    parser.add_argument("--skip-pictures", action="store_true")
    parser.add_argument("--skip-word-format", action="store_true")
    parser.add_argument("--skip-dictionary", action="store_true")
    parser.add_argument("--format", choices=["json", "markdown", "docx"], default="json")
    parser.add_argument("--out", help="Optional output file.")
    parser.add_argument("--report-dir", help="Optional directory for report.json, report.md, and report.docx.")
    args = parser.parse_args()

    path = Path(args.file).expanduser().resolve()
    result = check_file(path, args)
    if args.report_dir:
        result["report_paths"] = write_report_bundle(result, args.report_dir)

    if args.format == "json":
        text = json.dumps(result, ensure_ascii=False, indent=2)
    elif args.format == "markdown":
        text = render_markdown(result)
    else:
        if not args.out:
            raise SystemExit("--out is required when --format docx is used.")
        write_docx_report(result, args.out)
        return

    if args.out:
        Path(args.out).write_text(text, encoding="utf-8")
    else:
        print(text)


if __name__ == "__main__":
    main()
