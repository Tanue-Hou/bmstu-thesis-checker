"""Frozen counterfactual inputs for margin, table/image DOCX and embedded-font PDF."""
import hashlib,json,shutil,zipfile
from pathlib import Path
from io import BytesIO
from lxml import etree
from docx import Document
from docx.shared import Mm,Pt,RGBColor,Cm
from docx.enum.table import WD_ROW_HEIGHT_RULE
from PIL import Image,ImageDraw
import fitz
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'evidence/format-expansion'
STAGE=Path('C:/Users/Tanue Hou/AppData/Local/TestVkrLab-20261014')
SOURCE=ROOT/'evidence/graduate-structure/g_phd_left25.docx'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def record(path,parameters,parent=None):
    shutil.copyfile(path,STAGE/path.name)
    r={'case':path.stem,'format':path.suffix[1:],'sha256':sha(path),'parameters':parameters}
    if parent:
        r['parent_input']=str(parent.relative_to(ROOT));r['parent_sha256']=sha(parent)
        with zipfile.ZipFile(parent) as a,zipfile.ZipFile(path) as b:
            r['modified_zip_parts']=[n for n in b.namelist() if n not in a.namelist() or b.read(n)!=a.read(n)]
    return r
def main():
    OUT.mkdir(parents=True,exist_ok=True)
    if (OUT/'inputs.json').exists():raise SystemExit('Refusing to overwrite frozen inputs')
    records=[]
    for mm in [24,24.9,25.1,26]:
        d=Document(SOURCE);d.sections[0].left_margin=Mm(mm);p=OUT/('margin_'+str(mm).replace('.','_')+'.docx');d.save(p)
        records.append(record(p,{'left_margin_mm_requested':mm,'left_margin_twips':d.sections[0].left_margin.twips},SOURCE))
    chart=Image.new('RGB',(1600,800),'white');draw=ImageDraw.Draw(chart)
    draw.line([(100,50),(100,700),(1500,700)],fill='black',width=3)
    draw.line([(100,650),(400,510),(700,400),(1000,220),(1450,110)],fill='blue',width=5)
    for name,font,dimensions in [('complex_clean',14,None),('complex_table7',7,None),('complex_image_above',14,(4001,4000))]:
        d=Document(SOURCE);d.add_paragraph('Таблица 1 — Результаты исследования')
        t=d.add_table(rows=3,cols=2)
        for row in t.rows:
            row.height=Mm(8);row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
            for cell in row.cells:
                p=cell.paragraphs[0];p.text='Результаты обработки данных и анализа исследования.'
                p.paragraph_format.line_spacing=1;p.paragraph_format.first_line_indent=Cm(0)
                for r in p.runs:r.font.size=Pt(font);r.font.color.rgb=RGBColor(0,0,0)
        stream=BytesIO();(Image.new('RGB',dimensions,'white') if dimensions else chart).save(stream,format='PNG');stream.seek(0)
        d.add_picture(stream,width=Cm(12));d.add_paragraph('Рисунок 1 — Результаты анализа')
        p=OUT/(name+'.docx');d.save(p);records.append(record(p,{'table_font_pt':font,'image_dimensions':dimensions or [1600,800]},SOURCE))
    # One variable per PDF; embedded Times New Roman includes ToUnicode mapping.
    for name,font,page_mm,color in [('pdf_clean',14,(210,297),(0,0,0)),('pdf_font7',7,(210,297),(0,0,0)),('pdf_font8',8,(210,297),(0,0,0)),('pdf_font18',18,(210,297),(0,0,0)),('pdf_font18_5',18.5,(210,297),(0,0,0)),('pdf_letter',14,(216,279),(0,0,0)),('pdf_red',14,(210,297),(1,0,0))]:
        pdf=fitz.open();page=pdf.new_page(width=page_mm[0]*72/25.4,height=page_mm[1]*72/25.4)
        page.insert_font(fontname='TNR',fontfile='C:/Windows/Fonts/times.ttf')
        lines=['НАУЧНО-КВАЛИФИКАЦИОННАЯ РАБОТА','ОГЛАВЛЕНИЕ','ВВЕДЕНИЕ','Методы обработки данных и анализа.','Результаты исследования представлены.','ЗАКЛЮЧЕНИЕ','СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ']
        for i,text in enumerate(lines):page.insert_text((72,85+i*45),text,fontname='TNR',fontsize=font,color=color)
        p=OUT/(name+'.pdf');pdf.save(p,garbage=4,deflate=True);pdf.close()
        records.append(record(p,{'font_pt':font,'page_mm':page_mm,'text_color_rgb':color,'embedded_font':'Windows Times New Roman'}))
    (OUT/'inputs.json').write_text(json.dumps(records,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    (OUT/'fixture-modifications.json').write_text(json.dumps([r for r in records if 'parent_input' in r],ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    doc=fitz.open(OUT/'pdf_clean.pdf');doc[0].get_pixmap(matrix=fitz.Matrix(1,1)).save(str(OUT/'pdf_clean.render.png'));doc.close()
    print(json.dumps({'cases':len(records),'docx':sum(r['format']=='docx' for r in records),'pdf':sum(r['format']=='pdf' for r in records)}))
if __name__=='__main__':main()
