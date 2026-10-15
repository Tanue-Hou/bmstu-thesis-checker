"""Anonymous adversarial DOCX fixtures; no real thesis content."""
import json,hashlib,shutil,argparse
from pathlib import Path
from docx import Document
from docx.shared import Pt,Mm,Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image,ImageDraw
from repair_document import repair_copy
ROOT=Path(__file__).resolve().parents[1]
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def main():
    parser=argparse.ArgumentParser();parser.add_argument('--staging-dir',help='Optional explicit native-test staging directory; defaults to no external copies');args=parser.parse_args()
    out=ROOT/'evidence/v028/complex-repair';out.mkdir(parents=True,exist_ok=False)
    image=Image.new('RGB',(600,250),'white');d=ImageDraw.Draw(image);d.line((40,210,560,210),fill='black',width=2);d.line((40,210,40,20),fill='black',width=2);d.line((40,190,150,160,250,170,350,70,560,40),fill='blue',width=3);image.save(out/'plot.png')
    rows=[]
    for id,size,protected_only in [('C01',7,False),('C02',18.5,False),('C03',7,True)]:
        doc=Document();doc.styles['Normal'].font.size=Pt(14);doc.styles['Normal'].font.name='Times New Roman'
        sec=doc.sections[0];sec.page_width=Mm(210);sec.page_height=Mm(297)
        doc.add_paragraph('Министерство науки и высшего образования Российской Федерации')
        doc.add_paragraph('ВЫПУСКНАЯ КВАЛИФИКАЦИОННАЯ РАБОТА').runs[0].font.size=Pt(20)
        doc.add_paragraph('РЕФЕРАТ');doc.add_paragraph('Анонимный контроль сложного документа. '*15)
        doc.add_paragraph('СОДЕРЖАНИЕ');doc.add_paragraph('ВВЕДЕНИЕ 3');doc.add_page_break();doc.add_paragraph('ВВЕДЕНИЕ')
        p=doc.add_paragraph('Контроль основного текста и точечного исправления размера символов. '*5);p.runs[0].font.size=Pt(14 if protected_only else size)
        doc.add_page_break();doc.add_paragraph('Формулы и графические материалы',style='Heading 1')
        p=doc.add_paragraph('Уравнение сохраняется без изменения форматирования. '*3)
        math=OxmlElement('m:oMath');r=OxmlElement('m:r');t=OxmlElement('m:t');t.text='F = m a';r.append(t);math.append(r);p._p.append(math);p.runs[0].font.size=Pt(7)
        p=doc.add_paragraph('x');p.runs[0].font.size=Pt(6);p.runs[0].font.subscript=True
        doc.add_picture(str(out/'plot.png'),width=Inches(4));doc.add_paragraph('Рисунок 1 — Анонимный график',style='Caption')
        table=doc.add_table(rows=3,cols=3);table.style='Table Grid'
        for i,row in enumerate(table.rows):
            for j,cell in enumerate(row.cells):cell.text=f'{i+1}.{j+1}'
        p=doc.add_paragraph();hyper=OxmlElement('w:hyperlink');run=OxmlElement('w:r');text=OxmlElement('w:t');text.text='Protected hyperlink text';run.append(text);hyper.append(run);p._p.append(hyper)
        p=doc.add_paragraph();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),' PAGE ');p._p.append(field)
        control=OxmlElement('w:sdt');content=OxmlElement('w:sdtContent');control.append(content);p=OxmlElement('w:p');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='Protected content control';r.append(t);p.append(r);content.append(p);doc._element.body.insert(len(doc._element.body)-1,control)
        doc.sections[0].header.paragraphs[0].text='Protected header';doc.sections[0].footer.paragraphs[0].text='Protected footer'
        doc.add_paragraph('ЗАКЛЮЧЕНИЕ');doc.add_paragraph('Контрольный вывод. '*30);doc.add_paragraph('СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ');doc.add_paragraph('1. Анонимный источник.');doc.add_paragraph('ПРИЛОЖЕНИЕ А')
        before=out/(id+'B.docx');doc.save(before)
        result=repair_copy(before,out,actions=['font_size_anomalies'],body_font_pt=14,work_type='master',native_build=230)
        after=out/(id+'A.docx');shutil.copyfile(result['output_file'],after)
        if args.staging_dir:
            staging=Path(args.staging_dir).resolve();staging.mkdir(parents=True,exist_ok=True)
            for f in (before,after):shutil.copyfile(f,staging/f.name)
        rows.append({'id':id,'before':before.name,'after':after.name,'before_sha256':sha(before),'after_sha256':sha(after),'changes':len(result['changes']),'negative_protected_only':protected_only,'experimental_generator_build':230})
    (out/'inputs.json').write_text(json.dumps(rows,indent=2),encoding='utf-8');print(json.dumps(rows))
if __name__=='__main__':main()
