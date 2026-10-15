"""Conservative DOCX copy repair. Only the main document XML part is rewritten."""
import argparse
import hashlib
import json
import os
import sys
import uuid
import zipfile
from argparse import Namespace
from datetime import datetime, timezone
from pathlib import Path
from lxml import etree
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts/testvkr-service'))
from check_document import check_file, load_rule_registry, verification_context, context_evidence_valid
from parsers.docx_parser import DocxParser
from config import settings

ACTIONS = {'page_size_a4': ['Word.PageSizeNotA4'],
           'font_size_anomalies': ['Word.FontSizeTooSmall','Word.FontSizeTooLarge'],
           'text_color_black': ['Word.TextColorNotBlack']}

def sha(data):
    return hashlib.sha256(data).hexdigest()

def repair_copy(file_path, output_dir=None, actions=None, body_font_pt=None, work_type='bachelor', dry_run=False, expected_input_sha256=None, native_build=230):
    if native_build!=230:raise ValueError('Copy repair is validated against build230 only; build232 needs independent before/after evidence')
    source = Path(file_path).expanduser().resolve()
    if source.suffix.lower() != '.docx':
        raise ValueError('Automatic repair supports DOCX only; PDF remains read-only.')
    raw = source.read_bytes()
    if expected_input_sha256 and sha(raw)!=expected_input_sha256:raise ValueError("Source changed since preview; regenerate the preview")
    actions = list(ACTIONS) if actions is None else list(actions)
    if not all(a in ACTIONS for a in actions):
        raise ValueError('Unknown or unsupported repair action')
    registry = load_rule_registry()
    for action in actions:
        for rule_id in ACTIONS[action]:
            rule = registry.get(rule_id,{})
            evidence=verification_context(rule,'docx',work_type)
            if (rule.get('verification_status') != 'verified_against_original'
                or work_type not in rule.get('verified_work_types',[])
                or 'docx' not in rule.get('verified_formats',[])
                or evidence is None
                or any(getattr(settings,k,None)!=v for k,v in (evidence or {}).get('verified_policy',rule.get('verified_policy',{})).items())):
                raise ValueError('Repair requires validated policy/context: ' + rule_id)
            if work_type=='phd':
                repair_ref=evidence.get('repair_evidence_ref')
                if not repair_ref or sha((ROOT/repair_ref).read_bytes())!=evidence.get('repair_evidence_sha256'):
                    raise ValueError('Doctoral repair requires intact native before/after evidence: '+rule_id)
            refs=evidence.get('evidence_refs',[])
            if not context_evidence_valid(evidence):
                raise ValueError('Verification evidence is missing or changed: ' + rule_id)
    args=Namespace(work_type=work_type,skip_structure=False,skip_pictures=False,skip_word_format=False,skip_dictionary=False)
    before=check_file(source,args)
    triggered={f['rule_id'] for f in before['findings'] if f['verification_status']=='verified_against_original'}
    parser=DocxParser(str(source))
    document=parser.doc
    import re
    if re.search(r"(?:кандидата|доктора)\s+.{0,45}?наук", " ".join(p.text for p in document.paragraphs[:35]).lower()):raise ValueError("Dissertation-degree formatting repair requires independent validation; NKR proof is not transferable")
    if work_type=='phd' and (len(document.sections)!=1 or document.sections[0].page_width is None or document.sections[0].page_height is None or document.sections[0].page_width>document.sections[0].page_height):
        raise ValueError('Doctoral copy repair currently supports one portrait section only.')
    changes, skipped = [], []
    if 'page_size_a4' in actions and 'Word.PageSizeNotA4' in triggered:
        for i,sec in enumerate(document.sections):
            width,height=sec.page_width,sec.page_height
            if width is None or height is None or width > height:
                skipped.append({'section':i+1,'reason':'Unknown or landscape page; preserve layout'})
                continue
            if abs(width.mm-210)<=settings.PAGE_SIZE_TOLERANCE_MM and abs(height.mm-297)<=settings.PAGE_SIZE_TOLERANCE_MM:
                continue
            node=sec._sectPr.find(qn('w:pgSz'))
            old=dict(node.attrib)
            node.set(qn('w:w'),'11906');node.set(qn('w:h'),'16838')
            changes.append({'rule_id':'Word.PageSizeNotA4','section':i+1,'property':'page_size',
                            'before':old,'after':dict(node.attrib)})
    normal=document.styles['Normal'].font.size
    target=float(body_font_pt) if body_font_pt is not None else (normal.pt if normal is not None else None)
    if target is not None and not (settings.FONT_SIZE_MIN_PT <= target <= settings.FONT_SIZE_MAX_PT and target*2==int(target*2)):
        raise ValueError('Replacement body font must be a half-point value within validated anomaly limits')
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    cover_like=bool(re.search(r'министерство|квалификационн|пояснительная записка', ' '.join(p.text for p in document.paragraphs[:35]).lower()))
    intro=next((i for i,p in enumerate(document.paragraphs) if p.text.strip().lower()=='введение'),None)
    front_end=(intro if intro is not None else min(35,len(document.paragraphs))) if cover_like else 0
    for p_index,p in enumerate(document.paragraphs):
        if p_index<front_end:
            skipped.append({'paragraph':p_index+1,'reason':'Recognized cover/front matter; preserve intentional formatting'})
            continue
        if p._p.findall('.//w:fldChar',ns) or p._p.findall('.//w:fldSimple',ns) or p._p.findall('.//w:instrText',ns) or p._p.xpath('.//m:oMath'):
            skipped.append({'paragraph':p_index+1,'reason':'Field or equation paragraph'})
            continue
        for r_index,run in enumerate(p._p.findall('.//w:r',ns)):
            if run.xpath('ancestor::w:hyperlink | ancestor::w:ins | ancestor::w:del | ancestor::w:sdt | ancestor::w:txbxContent') or run.findall('.//w:drawing',ns) or run.findall('.//w:pict',ns):
                skipped.append({'paragraph':p_index+1,'run':r_index+1,'reason':'Hyperlink, drawing, text box, content control or tracked change'})
                continue
            props,_=parser.style_resolver.properties(p._p,run)
            text=''.join(n.text or '' for n in run.findall('.//w:t',ns))
            if not text.strip():continue
            size=props.find(qn('w:sz'))
            size_pt=float(size.get(qn('w:val')))/2 if size is not None else None
            color=props.find(qn('w:color'))
            color_hex=color.get(qn('w:val')) if color is not None else None
            modifications=[]
            if 'font_size_anomalies' in actions and size_pt is not None:
                rule_id='Word.FontSizeTooSmall' if size_pt<settings.FONT_SIZE_MIN_PT else ('Word.FontSizeTooLarge' if size_pt>settings.FONT_SIZE_MAX_PT else None)
                if rule_id in triggered:
                    if props.find(qn('w:vertAlign')) is not None or target is None or len(p.text.strip())<=40 or p.style.name.lower().startswith(('heading','title','caption')):
                        skipped.append({'paragraph':p_index+1,'run':r_index+1,'reason':'Special text or no known body font; leave font size unchanged'})
                    else:
                        modifications.append((rule_id,'sz',str(int(target*2))))
            if 'text_color_black' in actions and 'Word.TextColorNotBlack' in triggered and color_hex and color_hex.upper() not in ('000000','AUTO'):
                modifications.append(('Word.TextColorNotBlack','color','000000'))
            for rule_id,tag,value in modifications:
                rpr=run.find(qn('w:rPr'))
                if rpr is None:rpr=OxmlElement('w:rPr');run.insert(0,rpr)
                element=rpr.find(qn('w:'+tag))
                old=etree.tostring(element,encoding='unicode') if element is not None else None
                if element is None:element=OxmlElement('w:'+tag);rpr.append(element)
                if tag=='color':element.attrib.clear()
                element.set(qn('w:val'),value)
                changes.append({'rule_id':rule_id,'paragraph':p_index+1,'run':r_index+1,'xml_path':run.getroottree().getpath(run),
                                'text_excerpt':text[:100],'property':tag,'before_direct_xml':old,
                                'before_effective':size_pt if tag=='sz' else color_hex,
                                'after':target if tag=='sz' else value})
    table_manual={f['rule_id'] for f in before['findings'] if (f.get('threshold') or {}).get('table_count_chars',0)>0}
    for rule in sorted(table_manual):skipped.append({'rule_id':rule,'reason':'Table text is preserved; manual correction required.'})
    report={'schema_version':1,'created_at_utc':datetime.now(timezone.utc).isoformat(),
            'input_file':str(source),'input_sha256':sha(raw),'dry_run':dry_run,
            'native_build':native_build,'work_type':work_type,'actions':actions,'body_font_pt':target,'changes':changes,'skipped':skipped,
            'manual_only_rule_ids':sorted(({f['rule_id'] for f in before['findings']} - {r for a in actions for r in ACTIONS[a]})|table_manual),
            'scope':'Main document text and portrait sections only; no table/header/footer/image repair; original-program pass is not guaranteed.'}
    report["plan_id"]=sha(json.dumps({k:report[k] for k in ["input_sha256","work_type","actions","body_font_pt","changes","skipped"]},ensure_ascii=False,sort_keys=True).encode())
    if dry_run:
        report['source_unchanged']=source.read_bytes()==raw
        return report
    directory=Path(output_dir).expanduser().resolve() if output_dir else source.parent/'thesis_repairs'
    directory.mkdir(parents=True,exist_ok=True)
    destination=directory/(source.stem+'.repaired-'+uuid.uuid4().hex[:10]+'.docx')
    if destination.resolve()==source or destination.exists():raise ValueError('Output must be a new copy')
    with zipfile.ZipFile(source) as zin:
        entries=zin.infolist()
        if len({e.filename for e in entries})!=len(entries):raise ValueError('Duplicate ZIP members are unsupported')
        with destination.open('xb') as output, zipfile.ZipFile(output,'w') as zout:
            for info in entries:
                content=document.part.blob if info.filename=='word/document.xml' and changes else zin.read(info.filename)
                zout.writestr(info,content)
    after=check_file(destination,args)
    report.update({'output_file':str(destination),'output_sha256':sha(destination.read_bytes()),
                   'source_unchanged':source.read_bytes()==raw,
                   'before_rule_ids':[f['rule_id'] for f in before['findings']],
                   'after_rule_ids':[f['rule_id'] for f in after['findings']]})
    if not report['source_unchanged']:raise RuntimeError('Source changed during repair; do not accept this output')
    log_path=destination.with_suffix('.changes.json')
    with log_path.open('x',encoding='utf-8') as log:json.dump(report,log,ensure_ascii=False,indent=2)
    report['change_log']=str(log_path)
    return report

if __name__=='__main__':
    p=argparse.ArgumentParser()
    p.add_argument('file');p.add_argument('--output-dir');p.add_argument('--body-font-pt',type=float)
    p.add_argument('--actions',nargs='+',choices=list(ACTIONS));p.add_argument('--dry-run',action='store_true')
    p.add_argument('--work-type',choices=['bachelor','master','phd'],default='bachelor')
    p.add_argument('--expected-input-sha256')
    p.add_argument('--native-build',type=int,choices=[230],default=230)
    a=p.parse_args()
    result=repair_copy(a.file,a.output_dir,a.actions,a.body_font_pt,work_type=a.work_type,dry_run=a.dry_run,expected_input_sha256=a.expected_input_sha256,native_build=a.native_build)
    if hasattr(sys.stdout,'reconfigure'):sys.stdout.reconfigure(encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False,indent=2))
