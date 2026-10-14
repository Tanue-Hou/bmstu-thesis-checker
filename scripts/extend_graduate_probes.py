"""Reproduce the adaptive probes from frozen initial fixtures into a new directory."""
import argparse,json
from pathlib import Path
from docx import Document
from docx.shared import Cm
import generate_graduate_structure as gen

def main():
    p=argparse.ArgumentParser();p.add_argument('--output-dir',type=Path,required=True);a=p.parse_args()
    source=gen.OUT;destination=a.output_dir.resolve()
    if destination.exists():raise SystemExit('Choose a new reproduction directory')
    destination.mkdir(parents=True);gen.OUT=destination;gen.STAGE=destination/'stage';gen.STAGE.mkdir()
    volume=[]
    for page in [47,48]:
        d=Document(source/'g_pages48.docx')
        for para in d.paragraphs:
            if para.text.startswith('ПРИЛОЖЕНИЕ А ..........'):para.text='ПРИЛОЖЕНИЕ А .......... '+str(page)
        volume.append(gen.save(d,'g_toc'+str(page),48))
    d=Document(source/'g_phd_structure.docx')
    for para in d.paragraphs:
        if para.text=='СОДЕРЖАНИЕ':para.text='ОГЛАВЛЕНИЕ'
    phd=[gen.save(d,'g_phd_correct_heading')]
    dleft=Document(destination/'g_phd_correct_heading.docx');dleft.sections[0].left_margin=Cm(2.5)
    phd_margin=[gen.save(dleft,'g_phd_left25')]
    boundary=[];d=Document(destination/'g_phd_correct_heading.docx');d.sections[0].right_margin=Cm(1.5)
    boundary.append(gen.save(d,'g_phd_margin15'))
    for count in [5,10,11]:
        d=Document(source/'g_structure.docx');ps=d.paragraphs
        start=next(i for i,x in enumerate(ps) if x.text=='СОДЕРЖАНИЕ');end=next(i for i,x in enumerate(ps) if x.text=='ВВЕДЕНИЕ')
        rows=[x for x in ps[start+1:end] if x.text.strip()]
        for para in rows[count:]:para._p.getparent().remove(para._p)
        boundary.append(gen.save(d,'g_toc_entries'+str(count)))
    for target in [43199,43200]:
        d=Document(destination/'g_toc48.docx');anchor=next(x for x in d.paragraphs if x.text=='ЗАКЛЮЧЕНИЕ')
        for para in list(d.paragraphs):
            if para.text.startswith('Исследование методов обработки данных и анализ результатов.'):para._p.getparent().remove(para._p)
        ps=d.paragraphs;start=next(i for i,x in enumerate(ps) if x.text=='РЕФЕРАТ');end=next(i for i,x in enumerate(ps) if x.text=='ПРИЛОЖЕНИЕ А')
        remain=target-sum(len(x.text.strip()) for x in ps[start:end]);assert remain>0
        while remain>0:
            take=min(remain,600);text=('Методы исследования. '*40)[:take].rstrip();anchor.insert_paragraph_before(text+'а'*(take-len(text)));remain-=take
        boundary.append(gen.save(d,'g_chars'+str(target)))
    for name,records in [('volume-inputs',volume),('phd-inputs',phd),('boundary-inputs',boundary),('phd-margin-inputs',phd_margin)]:
        (destination/(name+'.json')).write_text(json.dumps(records,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'reproduced_cases':len(volume+phd+boundary+phd_margin),'output':str(destination)}))
if __name__=='__main__':main()
